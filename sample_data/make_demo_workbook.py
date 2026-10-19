"""Génère des classeurs de démonstration FICTIFS au format du classeur réel.

* ``demo_etats_financiers.xlsx`` : même structure que
  « Etats_Financiers_2021_2025.xlsx » (une feuille SYNTHÈSE, une feuille par
  ticker, ligne d'en-tête « Poste | 2021 ... 2025 », compte de résultat,
  BILAN, RATIOS, charges en négatif, unités hétérogènes).
* ``demo_base_marche.xlsx`` : même structure que « Base de données - SMF »
  (feuilles Cours, Nb_titres, Moyenne_cours) complétée d'une feuille
  Dividendes.

Les sociétés et les chiffres sont inventés (tickers préfixés « D ») : ils
servent à tester l'application sans connexion, jamais à décider.

Usage : python sample_data/make_demo_workbook.py [dossier_de_sortie]
"""
from __future__ import annotations

import math
import random
import sys
from datetime import date, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

YEARS = [2021, 2022, 2023, 2024, 2025]
HDR_FILL = PatternFill("solid", fgColor="1E2761")


def _sheet(wb, ticker, title, unit, blocks, years=YEARS):
    ws = wb.create_sheet(ticker)
    ws["A1"] = f"{title} ({ticker}) - Etats financiers (en {unit})"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Données fictives de démonstration"
    header = ["Poste"] + years + [f"Var {str(a)[2:]}/{str(b)[2:]}" for a, b in zip(years, years[1:])]
    ws.append([])
    ws.append(header)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HDR_FILL
    for block_title, rows in blocks:
        ws.append([block_title])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for label, vals in rows:
            r = ws.max_row + 1
            ws.append([label] + list(vals))
            for j in range(len(years) - 1):
                c1 = openpyxl.utils.get_column_letter(2 + j)
                c2 = openpyxl.utils.get_column_letter(3 + j)
                ws.cell(row=r, column=2 + len(years) + j,
                        value=f"=IF({c1}{r}=0,\"\",({c2}{r}-{c1}{r})/ABS({c1}{r}))")
    ws.column_dimensions["A"].width = 58
    return ws


def _grow(base, rates):
    out, v = [], base
    for r in rates:
        out.append(round(v, 1))
        v *= 1 + r
    return out


def build_financials(path: Path) -> None:
    wb = openpyxl.Workbook()
    syn = wb.active
    syn.title = "SYNTHÈSE"
    syn["A1"] = "Synthèse des sociétés (démonstration)"
    syn.append([])
    syn.append(["Société", "Secteur", "Unité"])
    synth_rows = []

    # ---------------- Banque (millions FCFA) ----------------
    pnb = _grow(98_000, [0.09, 0.11, 0.08, 0.10, 0])
    ip = [round(p * 1.05) for p in pnb]
    ic = [-round(p * 0.38) for p in pnb]
    cp_ = [round(p * 0.30) for p in pnb]
    cc = [-round(p * 0.03) for p in pnb]
    cge = [-round(p * r) for p, r in zip(pnb, [0.46, 0.45, 0.44, 0.43, 0.42])]
    da = [-round(p * 0.05) for p in pnb]
    rbe = [p + g + d for p, g, d in zip(pnb, cge, da)]
    cdr = [-round(p * r) for p, r in zip(pnb, [0.12, 0.10, 0.11, 0.09, 0.08])]
    re = [a + b for a, b in zip(rbe, cdr)]
    imp = [-round(r * 0.22) for r in re]
    rn = [a + b for a, b in zip(re, imp)]
    ta = _grow(1_950_000, [0.10, 0.09, 0.08, 0.07, 0])
    cred = [round(t * 0.58) for t in ta]
    dep = [round(t * 0.74) for t in ta]
    cpp = _grow(195_000, [0.11, 0.12, 0.10, 0.11, 0])
    _sheet(wb, "DBQC", "BANQUE DEMO COTE D'IVOIRE", "millions FCFA", [
        ("COMPTE DE RÉSULTAT", [
            ("Intérêts et produits assimilés", ip), ("Intérêts et charges assimilées", ic),
            ("Commissions (produits)", cp_), ("Commissions (charges)", cc),
            ("PRODUIT NET BANCAIRE", pnb), ("Charges générales d'exploitation", cge),
            ("Dotations aux amortissements & dépréciations", da),
            ("RÉSULTAT BRUT D'EXPLOITATION", rbe), ("Coût net du risque", cdr),
            ("RÉSULTAT D'EXPLOITATION", re), ("Impôts sur les bénéfices", imp), ("RÉSULTAT NET", rn)]),
        ("BILAN", []),
        ("ACTIF", [("Créances interbancaires", [round(t * 0.12) for t in ta]),
                   ("Créances sur la clientèle", cred),
                   ("Obligations et autres titres", [round(t * 0.22) for t in ta]),
                   ("TOTAL ACTIF", ta)]),
        ("PASSIF", [("Dettes interbancaires", [round(t * 0.08) for t in ta]),
                    ("Dettes à l'égard de la clientèle", dep),
                    ("Capital", [55_000] * 5),
                    ("Capitaux propres et ressources assimilées", cpp),
                    ("TOTAL PASSIF", ta)]),
        ("RATIOS", [("Coefficient d'exploitation", [round(-(g + d) / p * 100, 1) for g, d, p in zip(cge, da, pnb)]),
                    ("Taux de créances en souffrance (net)", [6.1, 5.8, 5.2, 4.9, 4.4]),
                    ("Dividende net par action", [330, 380, 420, 460, 500])]),
    ])
    synth_rows.append(("DBQC", "Banque", "M FCFA"))

    # ---------------- Banque (FCFA exacts) ----------------
    pnb2 = _grow(41_500e6, [0.06, 0.05, 0.07, 0.04, 0])
    cge2 = [-p * 0.55 for p in pnb2]
    da2 = [-p * 0.06 for p in pnb2]
    rbe2 = [a + b + c for a, b, c in zip(pnb2, cge2, da2)]
    cdr2 = [-p * r for p, r in zip(pnb2, [0.18, 0.22, 0.20, 0.25, 0.21])]
    rn2 = [(a + b) * 0.8 for a, b in zip(rbe2, cdr2)]
    ta2 = _grow(820_000e6, [0.05, 0.04, 0.06, 0.03, 0])
    _sheet(wb, "DBQS", "BANQUE DEMO SENEGAL", "FCFA", [
        ("COMPTE DE RÉSULTAT", [("PRODUIT NET BANCAIRE", pnb2), ("Charges générales d'exploitation", cge2),
                                ("Dotations aux amortissements & dépréciations", da2),
                                ("RÉSULTAT BRUT D'EXPLOITATION", rbe2), ("Coût net du risque", cdr2),
                                ("RÉSULTAT NET", rn2)]),
        ("BILAN", [("Créances sur la clientèle", [t * 0.62 for t in ta2]),
                   ("Dettes à l'égard de la clientèle", [t * 0.79 for t in ta2]),
                   ("Capitaux propres et ressources assimilées", _grow(78_000e6, [0.05, 0.05, 0.06, 0.05, 0])),
                   ("TOTAL ACTIF", ta2)]),
    ])
    synth_rows.append(("DBQS", "Banque", "FCFA"))

    # ---------------- Agro-industrie SYSCOHADA (millions) ----------------
    ca = _grow(182_000, [0.07, -0.04, 0.12, 0.06, 0])
    va = [round(c * r) for c, r in zip(ca, [0.31, 0.29, 0.33, 0.34, 0.35])]
    pers = [-round(c * 0.11) for c in ca]
    ebe = [a + b for a, b in zip(va, pers)]
    dot = [-round(c * 0.05) for c in ca]
    rexp = [a + b for a, b in zip(ebe, dot)]
    rfin = [-1_900, -2_100, -1_600, -1_200, -900]
    rao = [a + b for a, b in zip(rexp, rfin)]
    rhao = [300, -150, 220, 90, 400]
    imp3 = [-round((a + b) * 0.25) for a, b in zip(rao, rhao)]
    rn3 = [a + b + c for a, b, c in zip(rao, rhao, imp3)]
    tot = _grow(210_000, [0.04, 0.02, 0.06, 0.05, 0])
    cp3 = _grow(118_000, [0.06, 0.03, 0.08, 0.07, 0])
    _sheet(wb, "DAGC", "AGRO DEMO COTE D'IVOIRE", "millions FCFA", [
        ("COMPTE DE RÉSULTAT", [
            ("TA - Ventes de marchandises", [0] * 5), ("XB - Chiffre d'affaires (A + B + C)", ca),
            ("XC - VALEUR AJOUTÉE", va), ("RK - Charges de personnel", pers),
            ("XD - EXCÉDENT BRUT D'EXPLOITATION", ebe),
            ("RL - Dotations aux amortissements, aux provisions et dépréciations", dot),
            ("XE - RÉSULTAT D'EXPLOITATION", rexp), ("RM - Frais financiers et charges assimilées", [-2_500, -2_600, -2_100, -1_800, -1_500]),
            ("XF - RÉSULTAT FINANCIER", rfin), ("XG - RÉSULTAT DES ACTIVITÉS ORDINAIRES", rao),
            ("XH - RÉSULTAT HORS ACTIVITÉS ORDINAIRES", rhao), ("RS - Impôts sur le résultat", imp3),
            ("XI - RÉSULTAT NET", rn3)]),
        ("BILAN", []),
        ("ACTIF", [("ACTIF IMMOBILISÉ", [round(t * 0.55) for t in tot]), ("Stocks", [round(c * 0.12) for c in ca]),
                   ("Clients", [round(c * 0.10) for c in ca]),
                   ("ACTIF CIRCULANT", [round(t * 0.38) for t in tot]),
                   ("TRÉSORERIE-ACTIF", [round(t * 0.07) for t in tot]), ("TOTAL ACTIF", tot)]),
        ("PASSIF", [("Capital", [40_000] * 5), ("Réserves", [round(c * 0.4) for c in cp3]),
                    ("CAPITAUX PROPRES", cp3), ("DETTES FINANCIÈRES", [32_000, 30_000, 27_000, 22_000, 18_000]),
                    ("Fournisseurs", [round(c * 0.09) for c in ca]),
                    ("PASSIF CIRCULANT", [round(t * 0.30) for t in tot]),
                    ("TRÉSORERIE-PASSIF", [6_000, 7_500, 5_000, 4_000, 3_000]), ("TOTAL PASSIF", tot)]),
        ("RATIOS", [("Dividende net par action", [420, 480, 560, 620, 700])]),
    ])
    synth_rows.append(("DAGC", "Agro-industrie", "M FCFA"))

    # ---------------- Distribution SYSCOHADA (milliers), 2021-2022 saisis en FCFA ----
    ca4 = _grow(95_000_000, [0.05, 0.08, 0.04, 0.03, 0])
    mc4 = [round(c * 0.21) for c in ca4]
    va4 = [round(c * 0.14) for c in ca4]
    ebe4 = [round(v * 0.45) for v in va4]
    re4 = [round(e * 0.70) for e in ebe4]
    rn4 = [round(r * 0.66) for r in re4]
    tot4 = _grow(61_000_000, [0.03, 0.05, 0.02, 0.04, 0])
    cp4 = _grow(22_000_000, [0.04, 0.06, 0.05, 0.03, 0])
    rows4 = [("Chiffre d'affaires", ca4), ("Marge commerciale", mc4), ("Valeur ajoutée", va4),
             ("Excédent brut d'exploitation", ebe4), ("Résultat d'exploitation", re4), ("Résultat net", rn4)]
    bil4 = [("Total actif", tot4), ("Capitaux propres", cp4),
            ("Dettes financières", _grow(9_000_000, [0.1, 0.1, -0.2, -0.1, 0]))]

    def wrong_unit(v):  # erreur de saisie : deux exercices en FCFA exacts
        return [x * 1000 if i < 2 else x for i, x in enumerate(v)]
    _sheet(wb, "DDIC", "DISTRIBUTION DEMO COTE D'IVOIRE", "milliers FCFA", [
        ("COMPTE DE RÉSULTAT", [(l, wrong_unit(v)) for l, v in rows4]),
        ("BILAN", [(l, wrong_unit(v)) for l, v in bil4]),
    ])
    synth_rows.append(("DDIC", "Distribution", "K FCFA"))

    # ---------------- Télécom IFRS (millions) ----------------
    ca5 = _grow(1_480_000, [0.07, 0.08, 0.06, 0.06, 0])
    ebitda = [round(c * r) for c, r in zip(ca5, [0.47, 0.46, 0.46, 0.45, 0.45])]
    da5 = [-round(c * 0.15) for c in ca5]
    rop = [a + b for a, b in zip(ebitda, da5)]
    rn5 = [round(r * 0.70) for r in rop]
    rpg = [round(r * 0.84) for r in rn5]
    tot5 = _grow(2_400_000, [0.06, 0.07, 0.05, 0.05, 0])
    cpt = _grow(1_020_000, [0.04, 0.05, 0.05, 0.04, 0])
    cpg = [round(c * 0.83) for c in cpt]
    _sheet(wb, "DTLS", "TELECOM DEMO", "millions FCFA", [
        ("COMPTE DE RÉSULTAT CONSOLIDÉ", [("Chiffre d'affaires", ca5), ("EBITDA", ebitda),
                                           ("Dotations aux amortissements", da5),
                                           ("Résultat opérationnel", rop), ("Résultat net consolidé", rn5),
                                           ("Résultat net part du groupe", rpg)]),
        ("BILAN", [("Total actif", tot5), ("Capitaux propres part du groupe", cpg),
                   ("Total capitaux propres", cpt),
                   ("Dettes financières", _grow(420_000, [0.03, -0.02, 0.04, 0.01, 0])),
                   ("Trésorerie et équivalents de trésorerie", _grow(260_000, [0.05, 0.05, 0.05, 0.05, 0]))]),
        ("RATIOS", [("Dividende net par action", [1_700, 1_900, 2_050, 2_200, 2_400])]),
    ])
    synth_rows.append(("DTLS", "Télécom", "M FCFA"))

    # ---------------- Énergie SYSCOHADA (millions, marges en baisse) ----------------
    ca6 = _grow(410_000, [0.12, 0.02, -0.05, -0.03, 0])
    va6 = [round(c * r) for c, r in zip(ca6, [0.065, 0.06, 0.052, 0.046, 0.04])]
    ebe6 = [round(v * 0.5) for v in va6]
    re6 = [round(e * 0.7) for e in ebe6]
    rn6 = [round(r * 0.6) for r in re6]
    tot6 = _grow(160_000, [0.05, 0.04, 0.03, 0.03, 0])
    _sheet(wb, "DENC", "ENERGIE DEMO", "millions FCFA", [
        ("COMPTE DE RÉSULTAT", [("Chiffre d'affaires", ca6), ("Valeur ajoutée", va6),
                                ("Excédent brut d'exploitation", ebe6), ("Résultat d'exploitation", re6),
                                ("Résultat net", rn6)]),
        ("BILAN", [("Total actif", tot6), ("Capitaux propres", _grow(52_000, [0.03, 0.02, 0.0, -0.01, 0])),
                   ("Dettes financières", _grow(30_000, [0.1, 0.12, 0.15, 0.1, 0])),
                   ("Actif circulant", [round(t * 0.6) for t in tot6]),
                   ("Passif circulant", [round(t * 0.55) for t in tot6])]),
    ])
    synth_rows.append(("DENC", "Énergie", "M FCFA"))

    for r in synth_rows:
        syn.append(list(r))
    wb.save(path)


def build_market(path: Path, seed: int = 7) -> None:
    rnd = random.Random(seed)
    tickers = {
        # ticker: (cours final visé, dérive annuelle, volatilité annuelle, nb titres, flottant)
        "DBQC": (9_800, 0.18, 0.18, 60_000_000, 0.30),
        "DBQS": (2_250, 0.02, 0.20, 32_000_000, 0.25),
        "DAGC": (9_300, 0.12, 0.24, 25_000_000, 0.35),
        "DDIC": (2_550, -0.06, 0.26, 12_000_000, 0.40),
        "DTLS": (33_000, 0.10, 0.15, 100_000_000, 0.20),
        "DENC": (1_850, -0.10, 0.24, 40_000_000, 0.30),
        "BRVMC": (530.0, 0.25, 0.11, None, None),
    }
    start = date(2023, 1, 2)
    end = date(2026, 10, 16)
    days = []
    d = start
    while d <= end:
        if d.weekday() < 5:
            days.append(d)
        d += timedelta(days=1)
    prices = {}
    vols = {}
    for t, (p_end, mu, sig, *_r) in tickers.items():
        p = 100.0
        walk, vser = [], []
        for _ in days:
            ret = mu / 252 + sig / math.sqrt(252) * rnd.gauss(0, 1)
            ret = max(-0.075, min(0.075, ret))
            p = max(1.0, p * (1 + ret))
            walk.append(p)
        k = p_end / walk[-1]
        ser = [round(x * k, 2 if t == "BRVMC" else 0) for x in walk]
        for _ in days:
            traded = rnd.random() < (0.65 if t != "BRVMC" else 1)
            vser.append(0 if not traded else int(abs(rnd.gauss(4000, 3000))))
        prices[t] = ser
        vols[t] = vser

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cours"
    names = [t for t in tickers]
    ws.append(["N°", "Date"] + names)
    for i, day in enumerate(days):
        ws.append([i + 1, day] + [prices[t][i] for t in names])
    wv = wb.create_sheet("Volumes")
    wv.append(["N°", "Date"] + names[:-1])
    for i, day in enumerate(days):
        wv.append([i + 1, day] + [vols[t][i] for t in names[:-1]])
    wn = wb.create_sheet("Nb_titres")
    eq = [t for t in names if tickers[t][3]]
    wn.append(["", ""] + eq)
    wn.append(["Flottant", ""] + [int(tickers[t][3] * tickers[t][4]) for t in eq])
    wn.append(["Total", ""] + [tickers[t][3] for t in eq])
    wm = wb.create_sheet("Moyenne_cours")
    wm.append(["", "Année"] + names)
    for y in (2023, 2024, 2025, 2026):
        idx = [i for i, d in enumerate(days) if d.year == y]
        wm.append(["", y] + [round(sum(prices[t][i] for i in idx) / len(idx), 2) for t in names])
    wd = wb.create_sheet("Dividendes")
    wd.append(["Ticker", "Exercice", "Dividende net", "Date de paiement"])
    divs = {"DBQC": [380, 420, 460, 500], "DAGC": [480, 560, 620, 700], "DTLS": [1900, 2050, 2200, 2400],
            "DBQS": [120, 130, 0, 140], "DENC": [60, 55, 50, 45]}
    for t, ds in divs.items():
        for y, v in zip((2022, 2023, 2024, 2025), ds):
            wd.append([t, y, v, date(y + 1, 7, 15)])
    wb.save(path)


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent
    out.mkdir(parents=True, exist_ok=True)
    build_financials(out / "demo_etats_financiers.xlsx")
    build_market(out / "demo_base_marche.xlsx")
    print("Classeurs de démonstration générés dans", out)
