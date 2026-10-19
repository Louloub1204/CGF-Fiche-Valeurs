"""Fichier des soldes intermédiaires de gestion (SIG), suivi et renseigné par la société de gestion.

Troisième fichier de l'application, après le classeur des états financiers et la base de données SMF :

* ``build_template`` produit le classeur modèle (feuilles Lisez-moi, Entreprises, Banques, Titres, Suivi),
  pré-rempli avec ce que l'on sait déjà (classeur des états financiers, puis base SMF) ;
* ``parse_sig_file`` relit ce classeur une fois renseigné ;
* ``overlay`` fusionne les soldes lus dans la société analysée (priorité au fichier SIG pour un titre hors
  classeur ; simple complément des exercices manquants pour un titre du classeur).

Montants en millions de FCFA dans le fichier, convertis en FCFA à la lecture. Un exercice est une ligne.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field

from openpyxl import Workbook as XlWorkbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import config
from .utils import is_num, norm_label, to_float

NAVY, SKY, PAPER = "004B77", "EAF4FA", "F3F6F8"
YEARS = list(range(2021, 2026))
TRACK_YEARS = list(range(2021, 2027))
STATUTS = ["Publié", "Provisoire", "Estimé", "À vérifier"]
SHEETS = ("Entreprises", "Banques")
FIRST_ROW = 5
LAST_ROW = {"Entreprises": 420, "Banques": 260}
KEY_NONZERO = {"ca", "pnb", "total_actif", "capitaux_propres"}      # 0 = non renseigné pour ces postes

# (clé interne, intitulé de colonne, groupe)
ENTREPRISE_COLS = [
    ("ca", "Chiffre d'affaires", "Compte de résultat"),
    ("marge_commerciale", "Marge commerciale", "Compte de résultat"),
    ("va", "Valeur ajoutée", "Compte de résultat"),
    ("charges_personnel", "Charges de personnel", "Compte de résultat"),
    ("ebe", "Excédent brut d'exploitation (EBE)", "Compte de résultat"),
    ("dotations", "Dotations aux amortissements et provisions", "Compte de résultat"),
    ("re", "Résultat d'exploitation", "Compte de résultat"),
    ("rf", "Résultat financier", "Compte de résultat"),
    ("rao", "Résultat des activités ordinaires (RAO)", "Compte de résultat"),
    ("rhao", "Résultat hors activités ordinaires (HAO)", "Compte de résultat"),
    ("rai", "Résultat avant impôt", "Compte de résultat"),
    ("impots", "Impôts sur le résultat", "Compte de résultat"),
    ("rn", "Résultat net", "Compte de résultat"),
    ("rn_pg", "Résultat net part du groupe", "Compte de résultat"),
    ("capitaux_propres", "Capitaux propres", "Bilan"),
    ("total_actif", "Total du bilan", "Bilan"),
    ("dettes_financieres", "Dettes financières", "Bilan"),
]
BANQUE_COLS = [
    ("pnb", "Produit net bancaire (PNB)", "Compte de résultat"),
    ("marge_interets", "Marge d'intérêts", "Compte de résultat"),
    ("commissions_nettes", "Commissions nettes", "Compte de résultat"),
    ("charges_generales", "Charges générales d'exploitation", "Compte de résultat"),
    ("dotations", "Dotations aux amortissements", "Compte de résultat"),
    ("rbe", "Résultat brut d'exploitation (RBE)", "Compte de résultat"),
    ("cout_risque", "Coût du risque", "Compte de résultat"),
    ("re", "Résultat d'exploitation", "Compte de résultat"),
    ("impots", "Impôts sur les bénéfices", "Compte de résultat"),
    ("rn", "Résultat net", "Compte de résultat"),
    ("rn_pg", "Résultat net part du groupe", "Compte de résultat"),
    ("capitaux_propres", "Capitaux propres", "Bilan"),
    ("total_actif", "Total du bilan", "Bilan"),
    ("credits_clientele", "Crédits à la clientèle", "Bilan"),
    ("depots_clientele", "Dépôts de la clientèle", "Bilan"),
]
COLS = {"Entreprises": ENTREPRISE_COLS, "Banques": BANQUE_COLS}
FIXED = ["Ticker", "Société", "Exercice", "Statut", "Source"]           # colonnes A à E
N_FIXED = len(FIXED)


def sheet_of(ticker: str) -> str:
    ref = config.TICKERS.get(ticker, {})
    from .partial import kind_for
    return "Banques" if kind_for(ticker, ref.get("secteur", "")) == "banque" else "Entreprises"


def kind_label(ticker: str) -> str:
    from .partial import kind_for
    return {"banque": "Banque ou établissement financier", "ifrs": "Entreprise (normes IFRS)",
            "syscohada": "Entreprise (SYSCOHADA)"}[kind_for(ticker, config.TICKERS.get(ticker, {}).get("secteur", ""))]


# ---------------------------------------------------------------------------
# Lecture
# ---------------------------------------------------------------------------
@dataclass
class SigData:
    fields: dict[str, dict[str, dict[int, float]]] = field(default_factory=dict)   # ticker -> poste -> année -> FCFA
    estimated: set = field(default_factory=set)                                      # (ticker, poste, année)
    sources: dict[str, str] = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)
    unknown: list[str] = field(default_factory=list)
    n_rows: int = 0

    @property
    def tickers(self) -> list[str]:
        return sorted(t for t, f in self.fields.items() if f)

    def years(self, ticker: str) -> list[int]:
        ys = {y for s in self.fields.get(ticker, {}).values() for y in s}
        return sorted(ys)

    def count(self, ticker: str, name: str) -> int:
        return len(self.fields.get(ticker, {}).get(name, {}))


def _header_map(ws) -> tuple[int, dict[int, str]] | None:
    labels = {norm_label(lab): key for cols in COLS.values() for key, lab, _ in cols}
    for r in range(1, 12):
        if norm_label(ws.cell(r, 1).value) == "TICKER":
            m: dict[int, str] = {}
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if isinstance(v, str):
                    n = norm_label(v)
                    if n in labels:
                        m[c] = labels[n]
                    elif n == "EXERCICE":
                        m[c] = "_annee"
                    elif n == "STATUT":
                        m[c] = "_statut"
                    elif n == "SOURCE":
                        m[c] = "_source"
            return r, m
    return None


def _ticker(raw) -> str | None:
    if not isinstance(raw, str) or not raw.strip():
        return None
    t = raw.strip().upper()
    if t in config.TICKERS:
        return t
    from .market import _ticker_of
    cand = _ticker_of(t)
    return cand if cand in config.TICKERS else None


def parse_sig_file(source) -> SigData:
    """Classeur des soldes intermédiaires renseigné -> SigData (montants en FCFA)."""
    data = source.read() if hasattr(source, "read") else (bytes(source) if isinstance(source, (bytes, bytearray)) else None)
    wb = load_workbook(io.BytesIO(data) if data is not None else source, data_only=True, read_only=False)
    out = SigData()
    for name in wb.sheetnames:
        if norm_label(name) not in {norm_label(s) for s in SHEETS}:
            continue
        ws = wb[name]
        hm = _header_map(ws)
        if hm is None:
            out.notes.append(f"Feuille « {name} » ignorée : ligne d'en-tête (colonne « Ticker ») introuvable.")
            continue
        hr, cmap = hm
        c_t = next((c for c in range(1, 4) if norm_label(ws.cell(hr, c).value) == "TICKER"), 1)
        for r in range(hr + 1, ws.max_row + 1):
            t = _ticker(ws.cell(r, c_t).value)
            raw_t = ws.cell(r, c_t).value
            if t is None:
                if isinstance(raw_t, str) and raw_t.strip():
                    if raw_t.strip() not in out.unknown:
                        out.unknown.append(raw_t.strip())
                continue
            row = {key: ws.cell(r, c).value for c, key in cmap.items()}
            y = to_float(row.get("_annee"))
            if not is_num(y) or not 1990 < y < 2100:
                continue
            y = int(y)
            n_here = 0
            for key, v in row.items():
                if key.startswith("_"):
                    continue
                v = to_float(v)
                if not is_num(v) or (v == 0 and key in KEY_NONZERO):
                    continue
                out.fields.setdefault(t, {}).setdefault(key, {})[y] = float(v) * 1e6
                n_here += 1
                if norm_label(row.get("_statut")) == "ESTIME":
                    out.estimated.add((t, key, y))
            if n_here:
                out.n_rows += 1
                src = row.get("_source")
                if isinstance(src, str) and src.strip() and t not in out.sources:
                    out.sources[t] = src.strip()
    if not out.fields:
        out.notes.append("Aucun solde renseigné : vérifiez que les montants sont saisis sous les en-têtes du modèle "
                         "(feuilles « Entreprises » et « Banques », une ligne par titre et par exercice).")
    return out


# ---------------------------------------------------------------------------
# Fusion dans la société analysée
# ---------------------------------------------------------------------------
def overlay(c, sig: SigData | None, override: bool) -> int:
    """Fusionne les soldes du fichier SIG dans ``c.fields``. ``override`` : le fichier prime sur les valeurs
    existantes (titre hors classeur) ; sinon il ne comble que les exercices manquants (titre du classeur).
    Retourne le nombre de valeurs ajoutées ou remplacées."""
    fs = (sig.fields.get(c.ticker) if sig is not None else None) or {}
    if not fs:
        return 0
    n, conflicts = 0, []
    for name, series in fs.items():
        tgt = c.fields.setdefault(name, {})
        for y, v in series.items():
            old = tgt.get(y)
            if is_num(old) and not override:
                if old and abs(v - old) / abs(old) > 0.02:
                    conflicts.append(f"{c.field_labels.get(name, name)} {y}")
                continue
            tgt[y] = v
            n += 1
            if (c.ticker, name, y) in sig.estimated:
                c.estimated.add((name, y))
            else:
                c.estimated.discard((name, y))
    if conflicts:
        c.notes.append("Écart de plus de 2% entre le classeur des états financiers (retenu) et le fichier des "
                       "soldes intermédiaires pour : " + ", ".join(conflicts[:8]) + ("..." if len(conflicts) > 8 else "")
                       + ". Vérifiez la source.")
    if n:
        c.years = sorted(set(c.years) | {y for s in fs.values() for y in s})
        label = "Fichier des soldes intermédiaires"
        if label not in c.data_source:
            c.data_source = f"{label}, complété par {c.data_source}" if override else f"{c.data_source}, complété par le {label.lower()}"
        est = sorted({y for (t, k, y) in sig.estimated if t == c.ticker})
        if est:
            c.notes.append("Soldes marqués « Estimé » dans le fichier des soldes intermédiaires pour l'exercice "
                           + ", ".join(str(y) for y in est) + ".")
    return n


def level_after_overlay(c) -> str:
    """Niveau d'analyse d'un titre construit à partir du fichier SIG."""
    f = c.fields

    def n(name):
        return sum(1 for v in f.get(name, {}).values() if is_num(v))

    bilan = n(c.equity_field) >= 2 and n("total_actif") >= 2
    sig_n = sum(1 for k in ("va", "ebe", "re", "rao", "marge_interets", "rbe", "pnb") if n(k) >= 2 and k != c.revenue_field)
    if bilan and sig_n >= 2:
        return "complet"
    if bilan:
        return "partiel"
    return "chiffres_cles"


# ---------------------------------------------------------------------------
# Modèle Excel
# ---------------------------------------------------------------------------
def _prefill(tk: str, wb, mb) -> dict[int, tuple[dict[str, float], str]]:
    """Exercice -> ({poste: millions FCFA}, source) pour un titre, classeur d'abord, base SMF ensuite."""
    cols = [k for k, *_ in COLS[sheet_of(tk)]]
    rows: dict[int, tuple[dict[str, float], str]] = {}
    srcs: dict[int, set] = {}

    def put(y, key, v, src):
        if y not in YEARS or not is_num(v) or v == 0:
            return
        vals = rows.setdefault(y, ({}, ""))[0]
        if key not in vals:
            vals[key] = round(float(v) / 1e6, 3)
            srcs.setdefault(y, set()).add(src)

    if wb is not None and tk in getattr(wb, "companies", {}):
        c = wb.companies[tk]
        for key in cols:
            for y, v in c.fields.get(key, {}).items():
                put(y, key, v, "Classeur des états financiers")
    if mb is not None and tk in getattr(mb, "fundamentals", {}):
        f = mb.fundamentals[tk]
        is_bank = sheet_of(tk) == "Banques"
        rev = "pnb" if is_bank else "ca"
        for y, v in (f.get("pnb") or f.get("ca") or {}).items():
            put(y, rev, v, "Base SMF (Tableau de bord)")
        for key in ("rn", "capitaux_propres", "re"):
            if key in cols:
                for y, v in f.get(key, {}).items():
                    put(y, key, v, "Base SMF (Tableau de bord)")
    out = {}
    for y, (vals, _) in rows.items():
        s = srcs.get(y, set())
        out[y] = (vals, " + ".join(sorted(s, key=lambda x: (x != "Classeur des états financiers", x))))
    return out


def _tickers_by_sheet(extra: list[str] | None = None) -> dict[str, list[str]]:
    tk = list(config.TICKERS) + [t for t in (extra or []) if t not in config.TICKERS]
    out: dict[str, list[str]] = {"Entreprises": [], "Banques": []}
    for t in sorted(tk, key=lambda t: (config.TICKERS.get(t, {}).get("secteur", "~"), t)):
        out[sheet_of(t)].append(t)
    return out


def _col_letter(i: int) -> str:
    return get_column_letter(i)


def _check_formula(sheet: str, r: int, idx: dict[str, str]) -> str:
    """Contrôles de cohérence d'une ligne (texte vide si tout est cohérent)."""
    def c(k):
        return f"{idx[k]}{r}"

    parts = []
    if sheet == "Entreprises":
        parts.append(f'IF(COUNT({c("re")},{c("rf")},{c("rao")})=3,IF(ABS({c("re")}+{c("rf")}-{c("rao")})<=MAX(1,0.01*ABS({c("rao")})),"","RAO <> RE + RF "),"")')
        parts.append(f'IF(COUNT({c("rao")},{c("impots")},{c("rn")})=3,IF(ABS({c("rao")}+N({c("rhao")})-ABS({c("impots")})-{c("rn")})<=MAX(1,0.03*ABS({c("rn")})),"","RN <> RAO + HAO - impôt "),"")')
        parts.append(f'IF(COUNT({c("va")},{c("ebe")})=2,IF({c("ebe")}>{c("va")}+1,"EBE > VA ",""),"")')
        parts.append(f'IF(COUNT({c("ca")},{c("rn")})=2,IF(ABS({c("rn")})>1.5*ABS({c("ca")}),"RN > 150% du CA ",""),"")')
        parts.append(f'IF(COUNT({c("capitaux_propres")},{c("total_actif")})=2,IF({c("capitaux_propres")}>{c("total_actif")},"CP > total bilan ",""),"")')
    else:
        parts.append(f'IF(COUNT({c("rbe")},{c("cout_risque")},{c("re")})=3,IF(ABS({c("rbe")}+{c("cout_risque")}-{c("re")})<=MAX(1,0.015*ABS({c("re")})),"","RE <> RBE + coût du risque "),"")')
        parts.append(f'IF(COUNT({c("pnb")},{c("charges_generales")},{c("dotations")},{c("rbe")})=4,IF(ABS({c("pnb")}+{c("charges_generales")}+{c("dotations")}-{c("rbe")})<=MAX(1,0.015*ABS({c("rbe")})),"","RBE <> PNB + charges "),"")')
        parts.append(f'IF(COUNT({c("pnb")},{c("rn")})=2,IF(ABS({c("rn")})>{c("pnb")},"RN > PNB ",""),"")')
        parts.append(f'IF(COUNT({c("capitaux_propres")},{c("total_actif")})=2,IF({c("capitaux_propres")}>{c("total_actif")},"CP > total bilan ",""),"")')
    return "&".join(parts)


def build_template(wb=None, mb=None, years: list[int] | None = None, extra_tickers: list[str] | None = None) -> bytes:
    """Classeur modèle des soldes intermédiaires, pré-rempli depuis le classeur des états financiers (``wb``) et
    la base SMF (``mb``) lorsqu'ils sont fournis."""
    years = years or YEARS
    xl = XlWorkbook()
    thin = Side(style="thin", color="DCE4EA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    f_head = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    fill_head = PatternFill("solid", fgColor=NAVY)
    fill_bilan = PatternFill("solid", fgColor="2A6F9E")
    fill_fixed = PatternFill("solid", fgColor=PAPER)
    f_in = Font(name="Calibri", color="0000FF", size=10)        # saisies : bleu
    f_calc = Font(name="Calibri", color="000000", size=10)      # formules : noir
    f_base = Font(name="Calibri", size=10)
    num_fmt = '#,##0.0;-#,##0.0'

    # ---- Titres ----------------------------------------------------------
    ws_t = xl.active
    ws_t.title = "Titres"
    by_sheet = _tickers_by_sheet(extra_tickers)
    all_t = by_sheet["Entreprises"] + by_sheet["Banques"]
    ws_t.append(["Ticker", "Société", "Pays", "Secteur", "Type de comptes", "Feuille de saisie"])
    for t in all_t:
        ref = config.TICKERS.get(t, {})
        ws_t.append([t, ref.get("nom", t), config.COUNTRIES.get(ref.get("pays", ""), {}).get("nom", ref.get("pays", "")),
                     ref.get("secteur", ""), kind_label(t), sheet_of(t)])
    for c in ws_t[1]:
        c.font, c.fill = f_head, fill_head
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate([10, 38, 18, 30, 34, 16], 1):
        ws_t.column_dimensions[_col_letter(i)].width = w
    ws_t.freeze_panes = "A2"
    n_t = len(all_t) + 1
    for row in ws_t.iter_rows(min_row=2, max_row=n_t):
        for c in row:
            c.font, c.border = f_base, border

    # ---- Feuilles de saisie ------------------------------------------------
    idx_by_sheet: dict[str, dict[str, str]] = {}
    for sheet in SHEETS:
        ws = xl.create_sheet(sheet)
        cols = COLS[sheet]
        idx = {key: _col_letter(N_FIXED + 1 + i) for i, (key, _, _) in enumerate(cols)}
        idx_by_sheet[sheet] = idx
        c_comp = _col_letter(N_FIXED + len(cols) + 1)
        c_ctrl = _col_letter(N_FIXED + len(cols) + 2)
        last_c = N_FIXED + len(cols) + 2
        ws["A1"] = ("Soldes intermédiaires de gestion - " + ("entreprises (SYSCOHADA et IFRS)" if sheet == "Entreprises"
                                                              else "banques et établissements financiers"))
        ws["A1"].font = Font(name="Calibri", bold=True, size=14, color=NAVY)
        ws["A2"] = ("Une ligne par titre et par exercice. Montants en millions de FCFA, tels que publiés. Charges en négatif, "
                    "sauf les impôts sur le résultat (montant positif). Laissez vide ce que vous ne connaissez pas. "
                    "Texte en bleu : saisie. Voir la feuille « Lisez-moi ».")
        ws["A2"].font = Font(name="Calibri", italic=True, size=9, color="555555")
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=min(last_c, 12))
        ws.row_dimensions[2].height = 38
        # groupes (ligne 3) et en-têtes (ligne 4)
        ws.cell(3, 1, "Identification").font = Font(name="Calibri", bold=True, size=9, color=NAVY)
        grp_start = {}
        for i, (key, lab, grp) in enumerate(cols):
            grp_start.setdefault(grp, N_FIXED + 1 + i)
        for grp, c0 in grp_start.items():
            ws.cell(3, c0, grp).font = Font(name="Calibri", bold=True, size=9, color=NAVY)
        ws.cell(3, N_FIXED + len(cols) + 1, "Suivi").font = Font(name="Calibri", bold=True, size=9, color=NAVY)
        heads = FIXED + [lab for _, lab, _ in cols] + ["Complétude", "Contrôle de cohérence"]
        for j, h in enumerate(heads, 1):
            cell = ws.cell(4, j, h)
            cell.font = f_head
            cell.fill = fill_bilan if (1 <= j - N_FIXED <= len(cols) and cols[j - N_FIXED - 1][2] == "Bilan") else fill_head
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        ws.row_dimensions[4].height = 54
        widths = [10, 30, 9, 12, 28] + [15] * len(cols) + [11, 34]
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[_col_letter(j)].width = w
        ws.freeze_panes = ws.cell(FIRST_ROW, N_FIXED + 1)

        r = FIRST_ROW
        for t in by_sheet[sheet]:
            pre = _prefill(t, wb, mb)
            for y in years:
                vals, src = pre.get(y, ({}, ""))
                ws.cell(r, 1, t)
                ws.cell(r, 3, y)
                if vals:
                    ws.cell(r, 4, "Publié" if "Classeur" in src and "Base SMF" not in src else "À vérifier")
                    ws.cell(r, 5, src)
                for key, v in vals.items():
                    ws[f"{idx[key]}{r}"] = v
                r += 1
        last = LAST_ROW[sheet]
        if r > last:
            last = r + 50
            LAST_ROW[sheet] = last
        rng_c = f"{_col_letter(N_FIXED + 1)}{{r}}:{_col_letter(N_FIXED + len(cols))}{{r}}"
        for rr in range(FIRST_ROW, last + 1):
            ws.cell(rr, 2, f'=IF(A{rr}="","",IFERROR(VLOOKUP(A{rr},Titres!$A$2:$B$400,2,FALSE),"Ticker inconnu"))')
            ws.cell(rr, 2).font = f_calc
            ws.cell(rr, 2).fill = fill_fixed
            rg = rng_c.format(r=rr)
            ws[f"{c_comp}{rr}"] = f'=IF(A{rr}="","",COUNT({rg})/{len(cols)})'
            ws[f"{c_comp}{rr}"].number_format = "0%"
            chk = _check_formula(sheet, rr, idx)
            ws[f"{c_ctrl}{rr}"] = f'=IF(COUNT({rg})=0,"",IF(LEN({chk})=0,"OK",{chk}))'
            for cc in (c_comp, c_ctrl):
                ws[f"{cc}{rr}"].font = f_calc
            for j in range(1, last_c + 1):
                cell = ws.cell(rr, j)
                cell.border = border
                if j in (1, 3, 4, 5) or N_FIXED < j <= N_FIXED + len(cols):
                    cell.font = f_in
                if N_FIXED < j <= N_FIXED + len(cols):
                    cell.number_format = num_fmt
                    cell.alignment = Alignment(horizontal="right")
                if j == 3:
                    cell.alignment = Alignment(horizontal="center")
        # validations
        dv_t = DataValidation(type="list", formula1="=Titres!$A$2:$A$400", allow_blank=True, showErrorMessage=True,
                              errorTitle="Ticker", error="Choisissez un ticker de la feuille « Titres ».")
        dv_s = DataValidation(type="list", formula1='"' + ",".join(STATUTS) + '"', allow_blank=True)
        dv_y = DataValidation(type="whole", operator="between", formula1="1990", formula2="2100", allow_blank=True,
                              showErrorMessage=True, errorTitle="Exercice", error="Saisissez une année (ex. 2025).")
        dv_n = DataValidation(type="decimal", operator="between", formula1="-1E+12", formula2="1E+12", allow_blank=True,
                              showErrorMessage=True, errorTitle="Montant", error="Saisissez un nombre (millions de FCFA).")
        for dv, rg in ((dv_t, f"A{FIRST_ROW}:A{last}"), (dv_s, f"D{FIRST_ROW}:D{last}"), (dv_y, f"C{FIRST_ROW}:C{last}"),
                       (dv_n, f"{_col_letter(N_FIXED + 1)}{FIRST_ROW}:{_col_letter(N_FIXED + len(cols))}{last}")):
            ws.add_data_validation(dv)
            dv.add(rg)
        # mise en forme conditionnelle : contrôle en erreur, complétude
        ws.conditional_formatting.add(f"{c_ctrl}{FIRST_ROW}:{c_ctrl}{last}", FormulaRule(
            formula=[f'AND({c_ctrl}{FIRST_ROW}<>"",{c_ctrl}{FIRST_ROW}<>"OK")'],
            fill=PatternFill("solid", bgColor="F8D7DA"), font=Font(color="9B1C1C", bold=True)))
        ws.conditional_formatting.add(f"{c_ctrl}{FIRST_ROW}:{c_ctrl}{last}", FormulaRule(
            formula=[f'{c_ctrl}{FIRST_ROW}="OK"'], font=Font(color="1E7B34", bold=True)))
        ws.conditional_formatting.add(f"{c_comp}{FIRST_ROW}:{c_comp}{last}", ColorScaleRule(
            start_type="num", start_value=0, start_color="FFFFFF", end_type="num", end_value=1, end_color="9CCBE3"))
        ws.auto_filter.ref = f"A4:{_col_letter(last_c)}{last}"
        ws.sheet_view.zoomScale = 90

    # ---- Suivi -------------------------------------------------------------
    ws_s = xl.create_sheet("Suivi")
    ws_s["A1"] = "Suivi de la complétude des soldes intermédiaires"
    ws_s["A1"].font = Font(name="Calibri", bold=True, size=14, color=NAVY)
    ws_s["A2"] = ("Part des postes renseignés par titre et par exercice (calculée automatiquement à partir des feuilles "
                  "Entreprises et Banques). Un tiret signifie qu'aucun solde n'est saisi.")
    ws_s["A2"].font = Font(name="Calibri", italic=True, size=9, color="555555")
    hdr = ["Ticker", "Société", "Secteur"] + [str(y) for y in TRACK_YEARS] + ["Exercices renseignés"]
    for j, h in enumerate(hdr, 1):
        cell = ws_s.cell(4, j, h)
        cell.font, cell.fill, cell.border = f_head, fill_head, border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for j, y in enumerate(TRACK_YEARS):
        ws_s.cell(4, 4 + j, y)
    for i, t in enumerate(all_t):
        r = 5 + i
        ws_s.cell(r, 1, t).font = f_base
        ws_s.cell(r, 2, f'=IFERROR(VLOOKUP(A{r},Titres!$A$2:$D$400,2,FALSE),"")').font = f_base
        ws_s.cell(r, 3, f'=IFERROR(VLOOKUP(A{r},Titres!$A$2:$D$400,4,FALSE),"")').font = f_base
        for j, _y in enumerate(TRACK_YEARS):
            col = _col_letter(4 + j)
            terms = []
            for sheet in SHEETS:
                cols = COLS[sheet]
                c_comp = _col_letter(N_FIXED + len(cols) + 1)
                lr = LAST_ROW[sheet]
                terms.append(f"SUMIFS({sheet}!${c_comp}${FIRST_ROW}:${c_comp}${lr},{sheet}!$A${FIRST_ROW}:$A${lr},$A{r},"
                             f"{sheet}!$C${FIRST_ROW}:$C${lr},{col}$4)")
            ws_s[f"{col}{r}"] = "=" + "+".join(terms)
            ws_s[f"{col}{r}"].number_format = '0%;-0%;"-"'
            ws_s[f"{col}{r}"].alignment = Alignment(horizontal="center")
            ws_s[f"{col}{r}"].font = f_calc
        last_col = _col_letter(3 + len(TRACK_YEARS))
        ws_s[f"{_col_letter(4 + len(TRACK_YEARS))}{r}"] = f'=COUNTIF(D{r}:{last_col}{r},">0")'
        ws_s[f"{_col_letter(4 + len(TRACK_YEARS))}{r}"].alignment = Alignment(horizontal="center")
        for j in range(1, 5 + len(TRACK_YEARS)):
            ws_s.cell(r, j).border = border
    last_r = 4 + len(all_t)
    ws_s.conditional_formatting.add(f"D5:{_col_letter(3 + len(TRACK_YEARS))}{last_r}", ColorScaleRule(
        start_type="num", start_value=0, start_color="FFFFFF", end_type="num", end_value=1, end_color="6FB3D6"))
    for j, w in enumerate([10, 38, 30] + [10] * len(TRACK_YEARS) + [14], 1):
        ws_s.column_dimensions[_col_letter(j)].width = w
    ws_s.freeze_panes = "D5"

    # ---- Lisez-moi ----------------------------------------------------------
    ws_l = xl.create_sheet("Lisez-moi", 0)
    ws_l.sheet_view.showGridLines = False
    ws_l.column_dimensions["A"].width = 3
    ws_l.column_dimensions["B"].width = 30
    ws_l.column_dimensions["C"].width = 100
    ws_l["B2"] = "Soldes intermédiaires de gestion - CGF Gestion"
    ws_l["B2"].font = Font(name="Calibri", bold=True, size=18, color=NAVY)
    ws_l["B3"] = "Fichier de suivi à charger dans l'application « Fiche de valeur BRVM », au même titre que le classeur des états financiers et la base de données SMF."
    ws_l["B3"].font = Font(name="Calibri", italic=True, size=10, color="555555")
    lines = [
        ("Objet", "Centraliser, titre par titre et exercice par exercice, les soldes intermédiaires de gestion et les postes clés du bilan "
                  "que ni le classeur des états financiers ni la base SMF ne fournissent (valeur ajoutée, EBE, résultat d'exploitation, "
                  "RAO, total du bilan, etc.). L'application les utilise pour passer chaque titre à un niveau d'analyse plus complet."),
        ("Comment renseigner", "1. Ouvrez la feuille « Entreprises » ou « Banques » (selon le type du titre : voir la feuille « Titres »).\n"
                               "2. Repérez la ligne du titre et de l'exercice, ou ajoutez une ligne sous les dernières (les formules sont déjà prolongées).\n"
                               "3. Saisissez les montants en millions de FCFA, tels que publiés dans les états financiers ou le rapport annuel.\n"
                               "4. Indiquez le statut (Publié, Provisoire, Estimé, À vérifier) et la source (rapport annuel 2025, avis BRVM du...).\n"
                               "5. Contrôlez la colonne « Contrôle de cohérence » : elle signale les écarts de calcul entre soldes.\n"
                               "6. Suivez l'avancement dans la feuille « Suivi »."),
        ("Unité et signes", "Millions de FCFA. Les charges se saisissent en négatif (charges de personnel, dotations, charges générales "
                            "d'exploitation, coût du risque). Les impôts sur le résultat se saisissent en positif. Les soldes sont repris tels que publiés."),
        ("Statuts", "Publié : donnée issue des états financiers publiés. Provisoire : chiffres non audités ou communiqué de résultats. "
                    "Estimé : reconstitution par vos soins (l'application l'indique dans la fiche). À vérifier : valeur reprise de la base SMF ou du "
                    "classeur, à confirmer."),
        ("Lignes pré-remplies", "Les lignes déjà complétées proviennent du classeur des états financiers (statut « Publié ») ou de la base SMF "
                                "(statut « À vérifier »). Corrigez-les librement : le fichier prime sur la base SMF et les sources en ligne."),
        ("Règles de priorité", "Titre présent dans le classeur des états financiers : le classeur prime, le fichier SIG ne comble que les exercices "
                               "manquants et signale tout écart de plus de 2%. Titre absent du classeur : le fichier SIG prime sur la base SMF, "
                               "elle-même prioritaire sur les sources en ligne. Une saisie directe dans l'application reste prioritaire sur tout."),
        ("Niveau d'analyse obtenu", "Avec capitaux propres et total du bilan sur au moins deux exercices et deux soldes intermédiaires "
                                    "(valeur ajoutée, EBE, résultat d'exploitation, RAO ou RBE), le titre est analysé au niveau 1 (complet). "
                                    "Avec le bilan seulement, au niveau 1b. Sans bilan, il reste au niveau 2 (chiffres clés)."),
        ("Nouvel exercice", "Ajoutez simplement des lignes pour l'exercice suivant (par exemple 2026). Pour l'exercice en cours, la dernière "
                            "publication trimestrielle ou semestrielle est lue séparément par l'application."),
        ("Où trouver les chiffres", "États financiers annuels et rapports d'activités : brvm.org (avis et publications officielles), richbourse.com, "
                                    "sites des sociétés. Banques : PNB, coût du risque et résultat brut d'exploitation figurent au compte de résultat."),
        ("Définitions", "Marge commerciale = ventes de marchandises - achats de marchandises (variation de stocks incluse). "
                        "Valeur ajoutée = marge commerciale + production - consommations intermédiaires. "
                        "EBE = valeur ajoutée - impôts et taxes - charges de personnel. "
                        "Résultat d'exploitation = EBE - dotations + reprises + autres produits - autres charges. "
                        "RAO = résultat d'exploitation + résultat financier. Résultat net = RAO + HAO - impôts. "
                        "Banques : PNB = marge d'intérêts + commissions nettes + autres produits ; RBE = PNB - charges générales - dotations ; "
                        "résultat d'exploitation = RBE + coût du risque."),
        ("Code couleur", "Texte bleu : saisie ou donnée reprise d'une source. Texte noir : formule, ne pas modifier. Fond gris : colonne calculée. "
                         "En-tête bleu plus clair : postes du bilan."),
    ]
    r = 5
    for lab, txt in lines:
        ws_l.cell(r, 2, lab).font = Font(name="Calibri", bold=True, size=11, color=NAVY)
        ws_l.cell(r, 2).alignment = Alignment(vertical="top")
        ws_l.cell(r, 3, txt).font = Font(name="Calibri", size=10)
        ws_l.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")
        ws_l.row_dimensions[r].height = max(30, 15 * (txt.count("\n") + 1 + len(txt) // 105))
        for j in (2, 3):
            ws_l.cell(r, j).border = Border(bottom=thin)
        r += 1
    ws_l.cell(r + 1, 2, "Société de gestion collective au capital de 500 000 000 FCFA - N° agrément CREPMF SG-003/2001").font = \
        Font(name="Calibri", size=8, color="777777")

    order = ["Lisez-moi", "Entreprises", "Banques", "Suivi", "Titres"]
    xl._sheets = [xl[n] for n in order]
    xl.active = 0
    buf = io.BytesIO()
    xl.save(buf)
    return buf.getvalue()


def main(argv=None) -> int:
    import argparse
    from pathlib import Path
    ap = argparse.ArgumentParser(description="Génère le classeur modèle des soldes intermédiaires de gestion")
    ap.add_argument("--ef", help="Classeur des états financiers (pré-remplissage)")
    ap.add_argument("--smf", help="Base de données SMF (pré-remplissage)")
    ap.add_argument("--out", default="Soldes_Intermediaires_CGF.xlsx")
    a = ap.parse_args(argv)
    wb = mb = None
    if a.ef:
        from .parser import parse_workbook
        wb = parse_workbook(a.ef)
    if a.smf:
        from .market import load_market_base
        mb = load_market_base(a.smf)
    Path(a.out).write_bytes(build_template(wb, mb))
    print(f"Modèle écrit : {a.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
