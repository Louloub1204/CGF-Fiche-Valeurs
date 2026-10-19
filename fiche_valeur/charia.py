"""Conformité Charia : screening des titres selon quatre standards (FTSE, DJIM, S&P, Malaysia).

Méthode identique à celle du projet SMF (fichier « CHARIA SCREENING MODEL », feuille « TEST SCREENING UMOA ») :

- filtre préalable : revenus portant intérêt / chiffre d'affaires total <= 5% ;
- un standard est validé lorsque tous ses ratios respectent leur seuil ;
- un titre est compatible s'il valide au moins 3 standards sur 4 ;
- exclusions de principe : banques et établissements financiers conventionnels (Riba), tabac, jeux de hasard,
  boissons alcoolisées.

Source des ratios : le fichier de screening chargé (prioritaire) ; à défaut, recalcul à partir des états
financiers du classeur (dette financière, créances, trésorerie, total actif, capitaux propres).
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field

from . import config
from .utils import is_num

MIN_STANDARDS = 3
SEUIL_ILLICITE = 0.05

# Seuils effectivement appliqués (mise en forme conditionnelle du fichier de screening CGF)
STANDARDS: dict[str, dict] = {
    "FTSE": {"base": "total actif", "ratios": [("Dette / actif", "J", 0.33), ("Créances / actif", "K", 0.50),
                                                ("Trésorerie / actif", "L", 0.48)]},
    "DJIM": {"base": "capitalisation moyenne 24 mois", "ratios": [("Dette / capitalisation", "N", 0.33),
                                                                   ("Créances / capitalisation", "O", 0.33),
                                                                   ("Trésorerie / capitalisation", "P", 0.33)]},
    "S&P": {"base": "capitalisation moyenne 36 mois", "ratios": [("Dette / capitalisation", "R", 0.33),
                                                                  ("Créances / capitalisation", "S", 0.49),
                                                                  ("Trésorerie / capitalisation", "T", 0.33)]},
    "Malaysia": {"base": "total actif", "ratios": [("Dette / actif", "V", 0.33), ("Trésorerie / actif", "W", 0.48)]},
}

BANK_TICKERS = {"SGBC", "SIBC", "NSBC", "ECOC", "BICC", "BOAB", "BOAS", "BOABF", "BOAM", "BOAC", "BOAN", "CBIBF",
                "ETI", "ETIT", "ORGT", "SAFC", "BICB", "BBGC"}
ILLICIT_TICKERS = {"STBC": "Industrie du tabac", "LNBB": "Jeux de hasard et loterie",
                   "SLBC": "Production et distribution de boissons alcoolisées"}

SHEET_NAME = "TEST SCREENING UMOA"
ALIAS = {"SOGB": "SOGC"}      # symbole du fichier de screening -> ticker du référentiel


@dataclass
class StdResult:
    name: str
    base: str
    ok: bool
    ratios: list[tuple[str, float | None, float, bool | None]]     # libellé, valeur, seuil, respecté


@dataclass
class CharResult:
    ticker: str
    statut: str                      # compatible | non_compatible | exclu | non_evaluable
    n_pass: int = 0
    filtre_ok: bool | None = None
    illicit_ratio: float | None = None
    standards: list[StdResult] = field(default_factory=list)
    raison: str = ""
    source: str = ""
    annee: int | None = None
    notes: list[str] = field(default_factory=list)

    @property
    def evaluable(self) -> bool:
        return self.statut in ("compatible", "non_compatible")

    @property
    def label(self) -> str:
        if self.statut == "exclu":
            return "Exclu : " + self.raison.split(" - ")[0].split(" (")[0].lower() if self.raison else "Exclu"
        if self.statut == "non_evaluable":
            return "Non évalué"
        if self.filtre_ok is False:
            return f"Non compatible (revenus d'intérêt > 5%), {self.n_pass}/4"
        return ("Compatible" if self.statut == "compatible" else "Non compatible") + f" ({self.n_pass}/4)"


@dataclass
class CharData:
    results: dict[str, CharResult] = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)
    filename: str = ""

    @property
    def tickers(self) -> list[str]:
        return sorted(self.results)

    @property
    def n_rows(self) -> int:
        return len(self.results)


def _std_results(v: dict, halal: bool) -> tuple[list[StdResult], int]:
    """Évalue les quatre standards ; une donnée manquante ne pénalise pas (convention SMF)."""
    out, n = [], 0
    for name, spec in STANDARDS.items():
        ratios, ok = [], True
        for lab, key, seuil in spec["ratios"]:
            x = v.get(key)
            res = None if x is None else x <= seuil
            ratios.append((lab, x, seuil, res))
            ok = ok and res is not False
        out.append(StdResult(name, spec["base"], ok, ratios))
        n += int(ok)
    return out, n


def _verdict(ticker: str, v: dict, y: float | None, source: str, annee: int | None = None) -> CharResult:
    halal = y is None or y <= SEUIL_ILLICITE
    stds, n = _std_results(v, halal)
    n_eff = n if halal else 0
    compat = halal and n >= MIN_STANDARDS
    return CharResult(ticker, "compatible" if compat else "non_compatible", n_eff, halal, y, stds,
                      source=source, annee=annee)


def exclusion(ticker: str, sector: str = "", is_bank: bool = False) -> CharResult | None:
    if ticker in ILLICIT_TICKERS:
        return CharResult(ticker, "exclu", raison=f"Secteur illicite - {ILLICIT_TICKERS[ticker]}", source="règle de principe")
    if is_bank or ticker in BANK_TICKERS or sector == "Services financiers":
        return CharResult(ticker, "exclu", raison="Banque ou établissement financier conventionnel - modèle Riba",
                          source="règle de principe")
    return None


def _num(v):
    if isinstance(v, bool) or v is None or isinstance(v, str):
        return None
    try:
        x = float(v)
    except (TypeError, ValueError):
        return None
    return x if x == x else None


def parse_charia_file(source) -> CharData:
    """Lit la feuille « TEST SCREENING UMOA » (colonnes J-L FTSE, N-P DJIM, R-T S&P, V-W Malaysia, Y filtre, AA ticker)."""
    import openpyxl
    from openpyxl.utils import column_index_from_string as ci

    data = CharData()
    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    wb = openpyxl.load_workbook(source, data_only=True)
    name = next((n for n in wb.sheetnames if n.strip().upper() == SHEET_NAME), None) or \
        next((n for n in wb.sheetnames if "SCREENING" in n.upper()), None)
    if name is None:
        data.notes.append(f"Feuille « {SHEET_NAME} » introuvable : fichier de screening Charia non reconnu.")
        return data
    ws = wb[name]
    cols = {k: ci(k) for k in ("AA", "Y", "J", "K", "L", "N", "O", "P", "R", "S", "T", "V", "W")}
    unknown = 0
    for r in range(1, ws.max_row + 1):
        t = ws.cell(r, cols["AA"]).value
        if not isinstance(t, str) or not t.strip():
            continue
        t = ALIAS.get(t.strip().upper(), t.strip().upper())
        if t not in config.TICKERS:
            if 2 <= len(t) <= 6 and t.isalnum() and t != "SYMBOLE":
                unknown += 1
            continue
        vals = {k: _num(ws.cell(r, c).value) for k, c in cols.items() if k not in ("AA", "Y")}
        y = _num(ws.cell(r, cols["Y"]).value)
        present = [x for x in vals.values() if x is not None]
        exc = exclusion(t, config.TICKERS[t].get("secteur", ""))
        if exc is not None:
            data.results[t] = exc
            continue
        if not present or all(x == 0 for x in present):
            data.results[t] = CharResult(t, "non_evaluable", source="fichier de screening",
                                         raison="Données du fichier absentes ou nulles (cases vides, #DIV/0! ou zéros) : "
                                                "renseigner dette, créances, trésorerie et total actif")
            continue
        res = _verdict(t, vals, y, "fichier de screening")
        data.results[t] = res
    ev = [r_ for r_ in data.results.values() if r_.evaluable]
    same = [r_ for r_ in ev if all(a_[1] == b_[1] == c_[1] for a_, b_, c_ in zip(r_.standards[0].ratios, r_.standards[1].ratios,
                                                                                  r_.standards[2].ratios))]
    if ev and len(same) >= 0.9 * len(ev):
        msg = ("Les ratios DJIM et S&P du fichier sont identiques à ceux de FTSE : ils sont rapportés au total actif "
               "(capitalisations moyennes fixées à 1 dans la feuille). Renseigner les capitalisations pour un screening "
               "conforme à ces deux standards.")
        data.notes.append(msg)
        for r_ in ev:
            r_.notes.append(msg)
    if unknown:
        data.notes.append(f"{unknown} ligne(s) du fichier portent un symbole absent du référentiel et sont ignorées.")
    if not data.results:
        data.notes.append("Aucune ligne exploitable dans le fichier de screening Charia.")
    return data


def screen_company(c) -> CharResult | None:
    """Recalcul depuis les états financiers du classeur (dernier exercice avec total actif)."""
    f = c.fields
    ys = [y for y, v in f.get("total_actif", {}).items() if is_num(v) and v]
    if not ys:
        return None
    y = max(ys)
    g = lambda k: (abs(f[k][y]) if k in f and is_num(f[k].get(y)) else None)      # noqa: E731
    actif, dette, crean, cash = g("total_actif"), g("dettes_financieres"), g("creances_clients"), g("tresorerie_actif")
    if dette is None and crean is None and cash is None:
        return None
    cp_hist = sorted((yy, abs(v)) for yy, v in f.get("capitaux_propres", {}).items() if is_num(v) and v and yy <= y)
    cps = [v for _, v in cp_hist]
    cap24 = sum(cps[-2:]) / len(cps[-2:]) if cps else None
    cap36 = sum(cps[-3:]) / len(cps[-3:]) if cps else None
    r = lambda n, d: (n / d if n is not None and d else None)                       # noqa: E731
    vals = {"J": r(dette, actif), "K": r(crean, actif), "L": r(cash, actif),
            "N": r(dette, cap24), "O": r(crean, cap24), "P": r(cash, cap24),
            "R": r(dette, cap36), "S": r(crean, cap36), "T": r(cash, cap36),
            "V": r(dette, actif), "W": r(cash, actif)}
    ip = g("interets_produits")
    ca = g("ca")
    illicit = ip / ca if ip is not None and ca else None
    res = _verdict(c.ticker, vals, illicit, "états financiers du classeur (recalcul)", y)
    res.notes.append("Ratios recalculés : capitalisation estimée par la moyenne des capitaux propres des "
                     f"{min(len(cps), 2)} et {min(len(cps), 3)} derniers exercices.")
    if illicit is None:
        res.notes.append("Revenus portant intérêt non identifiés dans les comptes : filtre à 5% non vérifié.")
    return res


def resolve(ticker: str, c, data: CharData | None) -> CharResult:
    """Résultat Charia d'un titre : exclusion de principe, puis fichier chargé, puis recalcul, sinon non évalué."""
    sector = config.TICKERS.get(ticker, {}).get("secteur", "") or getattr(c, "sector", "")
    exc = exclusion(ticker, sector, bool(getattr(c, "is_bank", False)))
    if exc is not None:
        return exc
    if data is not None and ticker in data.results and data.results[ticker].statut != "non_evaluable":
        return data.results[ticker]
    calc = screen_company(c) if c is not None and getattr(c, "level", "") in ("complet", "partiel") else None
    if calc is not None:
        if data is not None and ticker in data.results:
            calc.notes.append("Le fichier de screening ne donne pas de données exploitables pour ce titre.")
        return calc
    if data is not None and ticker in data.results:
        return data.results[ticker]
    return CharResult(ticker, "non_evaluable",
                      raison="Charger le fichier de screening Charia ou les états financiers du titre")


def summary_line(r: CharResult) -> str:
    """Phrase de synthèse (fiche, commentaire)."""
    if r.statut == "exclu":
        return f"Titre exclu du périmètre Charia : {r.raison.lower()}."
    if r.statut == "non_evaluable":
        return "Conformité Charia non évaluée : " + (r.raison[0].lower() + r.raison[1:] if r.raison else "données absentes") + "."
    if r.filtre_ok is False:
        return (f"Non compatible Charia : revenus portant intérêt de {r.illicit_ratio * 100:.1f}% du chiffre d'affaires "
                f"pour un seuil de 5%.".replace(".", ",", 1))
    names = ", ".join(s.name for s in r.standards if s.ok) or "aucun"
    fail = [s.name for s in r.standards if not s.ok]
    txt = (f"{'Compatible' if r.statut == 'compatible' else 'Non compatible'} Charia : {r.n_pass} standard{'s' if r.n_pass > 1 else ''} sur 4 "
           f"validé{'s' if r.n_pass > 1 else ''} ({names}) ; au moins {MIN_STANDARDS} sont requis.")
    if fail:
        txt += " Non validé : " + ", ".join(fail) + "."
    return txt


def rows_for_table(r: CharResult) -> list[list[str]]:
    """Lignes à plat standard / ratio / valeur / seuil / résultat pour les tableaux des rapports."""
    out = []
    for s in r.standards:
        for lab, x, seuil, ok in s.ratios:
            out.append([s.name, lab, "n.d." if x is None else f"{x * 100:.1f}%".replace(".", ","),
                        f"<= {seuil * 100:.0f}%", "n.d." if ok is None else ("Oui" if ok else "Non")])
    return out
