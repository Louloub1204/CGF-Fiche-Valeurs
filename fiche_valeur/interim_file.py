"""Fichier des publications intermédiaires (T1, S1, T3) : modèle Excel, lecture, génération.

Troisième fichier de l'application, après le classeur des états financiers et la base de données SMF. Une ligne
par titre, exercice et période ; montants cumulés depuis le 1er janvier, en millions de FCFA. Chaque
publication fournit la période de l'exercice N et celle de N-1 : on renseigne les deux lignes.

* ``build_template`` produit le classeur (Lisez-moi, Entreprises, Banques, Suivi, Titres), pré-rempli avec les
  publications connues ;
* ``parse_interim_file`` le relit ;
* ``python -m fiche_valeur.interim_file --data publications.json --out Publications.xlsx`` le régénère.
"""
from __future__ import annotations

import io
import json
from datetime import date, datetime

from openpyxl import Workbook as XlWorkbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import config
from .dynamics import BANQUE_KEYS, ENTREPRISE_KEYS, PERIOD_ORDER, InterimData, PeriodRow, normalise_period
from .utils import is_num, norm_label, to_float

NAVY, PAPER = "004B77", "F3F6F8"
YEARS = list(range(2021, 2027))
PERIODS = ["T1", "S1", "T3"]
STATUTS = ["Publié", "Provisoire", "À vérifier"]
SHEETS = ("Entreprises", "Banques")
FIXED = ["Ticker", "Société", "Exercice", "Période", "Statut", "Date de publication", "Source (lien ou référence)"]
N_FIXED = len(FIXED)
FIRST_ROW = 5
COLS = {"Entreprises": [(k, lab) for k, lab, _ in ENTREPRISE_KEYS], "Banques": [(k, lab) for k, lab, _ in BANQUE_KEYS]}
HEAD_LABELS = {"ca": "Chiffre d'affaires", "va": "Valeur ajoutée", "ebe": "EBE (excédent brut d'exploitation)",
               "re": "Résultat d'exploitation", "rao": "Résultat des activités ordinaires (RAO)", "rn": "Résultat net",
               "pnb": "Produit net bancaire (PNB)", "charges_generales": "Frais de gestion (charges générales)",
               "rbe": "Résultat brut d'exploitation (RBE)", "cout_risque": "Coût du risque",
               "rai": "Résultat avant impôt", "credits_clientele": "Crédits nets à la clientèle",
               "depots_clientele": "Dépôts de la clientèle"}


def sheet_of(ticker: str) -> str:
    from .partial import kind_for
    return "Banques" if kind_for(ticker, config.TICKERS.get(ticker, {}).get("secteur", "")) == "banque" else "Entreprises"


def kind_label(ticker: str) -> str:
    from .partial import kind_for
    return {"banque": "Banque ou établissement financier", "ifrs": "Entreprise (normes IFRS)",
            "syscohada": "Entreprise (SYSCOHADA)"}[kind_for(ticker, config.TICKERS.get(ticker, {}).get("secteur", ""))]


# ---------------------------------------------------------------------------
# Lecture
# ---------------------------------------------------------------------------
def _ticker(raw) -> str | None:
    if not isinstance(raw, str) or not raw.strip():
        return None
    t = raw.strip().upper()
    if t in config.TICKERS:
        return t
    from .market import _ticker_of
    cand = _ticker_of(t)
    return cand if cand in config.TICKERS else None


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def parse_interim_file(source) -> InterimData:
    """Classeur des publications intermédiaires renseigné -> InterimData (montants en FCFA)."""
    data = source.read() if hasattr(source, "read") else (bytes(source) if isinstance(source, (bytes, bytearray)) else None)
    wb = load_workbook(io.BytesIO(data) if data is not None else source, data_only=True, read_only=False)
    out = InterimData()
    labels = {norm_label(HEAD_LABELS[k]): k for cols in COLS.values() for k, _ in cols}
    meta = {norm_label(h): h for h in FIXED}
    for name in wb.sheetnames:
        if norm_label(name) not in {norm_label(s) for s in SHEETS}:
            continue
        ws = wb[name]
        hr = next((r for r in range(1, 12) if norm_label(ws.cell(r, 1).value) == "TICKER"), None)
        if hr is None:
            out.notes.append(f"Feuille « {name} » ignorée : ligne d'en-tête (colonne « Ticker ») introuvable.")
            continue
        cmap: dict[int, str] = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(hr, c).value
            if not isinstance(v, str):
                continue
            n = norm_label(v)
            if n in labels:
                cmap[c] = labels[n]
            elif n in meta:
                cmap[c] = "_" + meta[n]
        for r in range(hr + 1, ws.max_row + 1):
            row = {key: ws.cell(r, c).value for c, key in cmap.items()}
            raw_t = row.get("_Ticker")
            t = _ticker(raw_t)
            if t is None:
                if isinstance(raw_t, str) and raw_t.strip() and raw_t.strip() not in out.unknown:
                    out.unknown.append(raw_t.strip())
                continue
            y = to_float(row.get("_Exercice"))
            per = normalise_period(row.get("_Période"))
            if not is_num(y) or not 1990 < y < 2100 or per is None:
                continue
            vals = {}
            for key, v in row.items():
                if key.startswith("_"):
                    continue
                v = to_float(v)
                if is_num(v):
                    vals[key] = float(v) * 1e6
            if not vals:
                continue
            statut = row.get("_Statut")
            out.rows.setdefault(t, {})[(int(y), per)] = PeriodRow(
                values=vals, statut=str(statut).strip() if statut else "", published=_date(row.get("_Date de publication")),
                source=str(row.get("_Source (lien ou référence)") or "").strip(), origin="fichier")
            out.n_rows += 1
    if not out.rows:
        out.notes.append("Aucune publication renseignée : saisissez les montants sous les en-têtes du modèle "
                         "(feuilles « Entreprises » et « Banques », une ligne par titre, exercice et période).")
    return out


# ---------------------------------------------------------------------------
# Modèle Excel
# ---------------------------------------------------------------------------
def _check(sheet: str, r: int, idx: dict[str, str]) -> str:
    def c(k):
        return f"{idx[k]}{r}"
    if sheet == "Entreprises":
        parts = [f'IF(COUNT({c("va")},{c("ebe")})=2,IF({c("ebe")}>{c("va")}+1,"EBE > VA ",""),"")',
                 f'IF(COUNT({c("ca")},{c("rn")})=2,IF(ABS({c("rn")})>ABS({c("ca")}),"RN > CA ",""),"")',
                 f'IF(COUNT({c("ca")},{c("re")})=2,IF(ABS({c("re")})>ABS({c("ca")}),"RE > CA ",""),"")']
    else:
        parts = [f'IF(COUNT({c("pnb")},{c("charges_generales")},{c("rbe")})=3,IF(ABS({c("pnb")}+{c("charges_generales")}-{c("rbe")})<=MAX(1,0.02*ABS({c("rbe")})),"","RBE <> PNB + frais "),"")',
                 f'IF(COUNT({c("rbe")},{c("cout_risque")},{c("rai")})=3,IF(ABS({c("rbe")}+{c("cout_risque")}-{c("rai")})<=MAX(1,0.03*ABS({c("rai")})),"","RAI <> RBE + coût du risque "),"")',
                 f'IF(COUNT({c("pnb")},{c("rn")})=2,IF(ABS({c("rn")})>ABS({c("pnb")}),"RN > PNB ",""),"")']
    return "&".join(parts)


def tickers_by_sheet() -> dict[str, list[str]]:
    out: dict[str, list[str]] = {"Entreprises": [], "Banques": []}
    for t in sorted(config.TICKERS, key=lambda t: (config.TICKERS[t].get("secteur", "~"), t)):
        out[sheet_of(t)].append(t)
    return out


def build_template(records: list[dict] | None = None, years: list[int] | None = None) -> bytes:
    """Classeur modèle des publications intermédiaires, pré-rempli avec ``records`` : liste de dictionnaires
    ``{ticker, year, period, statut, published, source, values: {clé: millions FCFA}}``."""
    years = years or YEARS
    pre = {(r["ticker"], int(r["year"]), r["period"]): r for r in (records or [])}
    xl = XlWorkbook()
    thin = Side(style="thin", color="DCE4EA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    f_head = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    fill_head = PatternFill("solid", fgColor=NAVY)
    fill_fixed = PatternFill("solid", fgColor=PAPER)
    f_in = Font(name="Calibri", color="0000FF", size=10)
    f_calc = Font(name="Calibri", color="000000", size=10)
    f_base = Font(name="Calibri", size=10)
    num_fmt = '#,##0.0;-#,##0.0'
    by_sheet = tickers_by_sheet()
    all_t = by_sheet["Entreprises"] + by_sheet["Banques"]

    ws_t = xl.active
    ws_t.title = "Titres"
    ws_t.append(["Ticker", "Société", "Pays", "Secteur", "Type de comptes", "Feuille de saisie"])
    for t in all_t:
        ref = config.TICKERS[t]
        ws_t.append([t, ref.get("nom", t), config.COUNTRIES.get(ref.get("pays", ""), {}).get("nom", ref.get("pays", "")),
                     ref.get("secteur", ""), kind_label(t), sheet_of(t)])
    for c in ws_t[1]:
        c.font, c.fill = f_head, fill_head
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate([10, 38, 18, 30, 34, 16], 1):
        ws_t.column_dimensions[get_column_letter(i)].width = w
    ws_t.freeze_panes = "A2"
    for row in ws_t.iter_rows(min_row=2, max_row=len(all_t) + 1):
        for c in row:
            c.font, c.border = f_base, border

    last_row: dict[str, int] = {}
    comp_col: dict[str, str] = {}
    for sheet in SHEETS:
        ws = xl.create_sheet(sheet)
        cols = COLS[sheet]
        idx = {k: get_column_letter(N_FIXED + 1 + i) for i, (k, _) in enumerate(cols)}
        c_comp = get_column_letter(N_FIXED + len(cols) + 1)
        c_ctrl = get_column_letter(N_FIXED + len(cols) + 2)
        last_c = N_FIXED + len(cols) + 2
        comp_col[sheet] = c_comp
        ws["A1"] = "Publications intermédiaires (T1, S1, T3) - " + ("entreprises" if sheet == "Entreprises" else "banques et établissements financiers")
        ws["A1"].font = Font(name="Calibri", bold=True, size=14, color=NAVY)
        ws["A2"] = ("Une ligne par titre, exercice et période. Montants cumulés depuis le 1er janvier (T1 = 3 mois, S1 = 6 mois, "
                    "T3 = 9 mois), en millions de FCFA, tels que publiés. Charges en négatif. Une publication donne l'exercice N et "
                    "le comparatif N-1 : renseignez les deux lignes. Texte bleu : saisie. Voir « Lisez-moi ».")
        ws["A2"].font = Font(name="Calibri", italic=True, size=9, color="555555")
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=min(last_c, 12))
        ws.row_dimensions[2].height = 40
        heads = FIXED + [HEAD_LABELS[k] for k, _ in cols] + ["Complétude", "Contrôle de cohérence"]
        for j, h in enumerate(heads, 1):
            cell = ws.cell(4, j, h)
            cell.font, cell.fill, cell.border = f_head, fill_head, border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[4].height = 54
        for j, w in enumerate([10, 30, 9, 9, 12, 14, 40] + [15] * len(cols) + [11, 30], 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = ws.cell(FIRST_ROW, N_FIXED + 1)
        r = FIRST_ROW
        for t in by_sheet[sheet]:
            for y in years:
                for p in PERIODS:
                    rec = pre.get((t, y, p))
                    ws.cell(r, 1, t)
                    ws.cell(r, 3, y)
                    ws.cell(r, 4, p)
                    if rec:
                        ws.cell(r, 5, rec.get("statut", "Publié"))
                        if rec.get("published"):
                            ws.cell(r, 6, _date(rec["published"]) or rec["published"])
                        ws.cell(r, 7, rec.get("source", ""))
                        for k, v in (rec.get("values") or {}).items():
                            if k in idx and v is not None:
                                ws[f"{idx[k]}{r}"] = v
                    r += 1
        last = r + 40                              # marge de lignes libres sous les titres
        last_row[sheet] = last
        rg_tpl = f"{get_column_letter(N_FIXED + 1)}{{r}}:{get_column_letter(N_FIXED + len(cols))}{{r}}"
        for rr in range(FIRST_ROW, last + 1):
            ws.cell(rr, 2, f'=IF(A{rr}="","",IFERROR(VLOOKUP(A{rr},Titres!$A$2:$B$400,2,FALSE),"Ticker inconnu"))')
            ws.cell(rr, 2).font, ws.cell(rr, 2).fill = f_calc, fill_fixed
            rg = rg_tpl.format(r=rr)
            ws[f"{c_comp}{rr}"] = f'=IF(A{rr}="","",COUNT({rg})/{len(cols)})'
            ws[f"{c_comp}{rr}"].number_format = "0%"
            chk = _check(sheet, rr, idx)
            ws[f"{c_ctrl}{rr}"] = f'=IF(COUNT({rg})=0,"",IF(LEN({chk})=0,"OK",{chk}))'
            for cc in (c_comp, c_ctrl):
                ws[f"{cc}{rr}"].font = f_calc
            for j in range(1, last_c + 1):
                cell = ws.cell(rr, j)
                cell.border = border
                if j in (1, 3, 4, 5, 6, 7) or N_FIXED < j <= N_FIXED + len(cols):
                    cell.font = f_in
                if N_FIXED < j <= N_FIXED + len(cols):
                    cell.number_format = num_fmt
                    cell.alignment = Alignment(horizontal="right")
                if j in (3, 4):
                    cell.alignment = Alignment(horizontal="center")
                if j == 6:
                    cell.number_format = "dd/mm/yyyy"
        dvs = [(DataValidation(type="list", formula1="=Titres!$A$2:$A$400", allow_blank=True, showErrorMessage=True,
                               errorTitle="Ticker", error="Choisissez un ticker de la feuille « Titres »."), f"A{FIRST_ROW}:A{last}"),
               (DataValidation(type="whole", operator="between", formula1="1990", formula2="2100", allow_blank=True,
                               showErrorMessage=True, errorTitle="Exercice", error="Saisissez une année (ex. 2026)."), f"C{FIRST_ROW}:C{last}"),
               (DataValidation(type="list", formula1='"T1,S1,T3"', allow_blank=True, showErrorMessage=True,
                               errorTitle="Période", error="T1, S1 ou T3."), f"D{FIRST_ROW}:D{last}"),
               (DataValidation(type="list", formula1='"' + ",".join(STATUTS) + '"', allow_blank=True), f"E{FIRST_ROW}:E{last}"),
               (DataValidation(type="decimal", operator="between", formula1="-1E+12", formula2="1E+12", allow_blank=True,
                               showErrorMessage=True, errorTitle="Montant", error="Saisissez un nombre (millions de FCFA)."),
                f"{get_column_letter(N_FIXED + 1)}{FIRST_ROW}:{get_column_letter(N_FIXED + len(cols))}{last}")]
        for dv, rg in dvs:
            ws.add_data_validation(dv)
            dv.add(rg)
        ws.conditional_formatting.add(f"{c_ctrl}{FIRST_ROW}:{c_ctrl}{last}", FormulaRule(
            formula=[f'AND({c_ctrl}{FIRST_ROW}<>"",{c_ctrl}{FIRST_ROW}<>"OK")'],
            fill=PatternFill("solid", bgColor="F8D7DA"), font=Font(color="9B1C1C", bold=True)))
        ws.conditional_formatting.add(f"{c_ctrl}{FIRST_ROW}:{c_ctrl}{last}", FormulaRule(
            formula=[f'{c_ctrl}{FIRST_ROW}="OK"'], font=Font(color="1E7B34", bold=True)))
        ws.conditional_formatting.add(f"{c_comp}{FIRST_ROW}:{c_comp}{last}", ColorScaleRule(
            start_type="num", start_value=0, start_color="FFFFFF", end_type="num", end_value=1, end_color="9CCBE3"))
        ws.auto_filter.ref = f"A4:{get_column_letter(last_c)}{last}"
        ws.sheet_view.zoomScale = 90

    # ---- Suivi : titres x (exercice, période) ------------------------------
    ws_s = xl.create_sheet("Suivi")
    ws_s["A1"] = "Suivi des publications intermédiaires"
    ws_s["A1"].font = Font(name="Calibri", bold=True, size=14, color=NAVY)
    ws_s["A2"] = ("Part des indicateurs renseignés par titre et par période (calculée à partir des feuilles Entreprises et Banques). "
                  "Un tiret signifie qu'aucune publication n'est saisie.")
    ws_s["A2"].font = Font(name="Calibri", italic=True, size=9, color="555555")
    ws_s["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws_s.merge_cells("A2:C3")
    combos = [(y, p) for y in years for p in PERIODS]
    hdr = ["Ticker", "Société", "Secteur"] + [f"{p} {y}" for y, p in combos]
    for j, h in enumerate(hdr, 1):
        cell = ws_s.cell(4, j, h)
        cell.font, cell.fill, cell.border = f_head, fill_head, border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for j, (y, p) in enumerate(combos):                        # lignes 3 (exercice) et en-têtes cachés pour les formules
        ws_s.cell(3, 4 + j, y).font = Font(color="FFFFFF", size=6)
        ws_s.cell(2, 4 + j, p).font = Font(color="FFFFFF", size=6)
    for i, t in enumerate(all_t):
        rr = 5 + i
        ws_s.cell(rr, 1, t).font = f_base
        ws_s.cell(rr, 2, f'=IFERROR(VLOOKUP(A{rr},Titres!$A$2:$D$400,2,FALSE),"")').font = f_base
        ws_s.cell(rr, 3, f'=IFERROR(VLOOKUP(A{rr},Titres!$A$2:$D$400,4,FALSE),"")').font = f_base
        for j, _ in enumerate(combos):
            col = get_column_letter(4 + j)
            terms = []
            for sheet in SHEETS:
                cc, lr = comp_col[sheet], last_row[sheet]
                terms.append(f"SUMIFS({sheet}!${cc}${FIRST_ROW}:${cc}${lr},{sheet}!$A${FIRST_ROW}:$A${lr},$A{rr},"
                             f"{sheet}!$C${FIRST_ROW}:$C${lr},{col}$3,{sheet}!$D${FIRST_ROW}:$D${lr},{col}$2)")
            ws_s[f"{col}{rr}"] = "=" + "+".join(terms)
            ws_s[f"{col}{rr}"].number_format = '0%;-0%;"-"'
            ws_s[f"{col}{rr}"].alignment = Alignment(horizontal="center")
            ws_s[f"{col}{rr}"].font = f_calc
        for j in range(1, 4 + len(combos)):
            ws_s.cell(rr, j).border = border
    last_r = 4 + len(all_t)
    ws_s.conditional_formatting.add(f"D5:{get_column_letter(3 + len(combos))}{last_r}", ColorScaleRule(
        start_type="num", start_value=0, start_color="FFFFFF", end_type="num", end_value=1, end_color="6FB3D6"))
    for j, w in enumerate([10, 36, 28] + [9] * len(combos), 1):
        ws_s.column_dimensions[get_column_letter(j)].width = w
    ws_s.freeze_panes = "D5"

    # ---- Lisez-moi ------------------------------------------------------------
    ws_l = xl.create_sheet("Lisez-moi", 0)
    ws_l.sheet_view.showGridLines = False
    ws_l.column_dimensions["A"].width = 3
    ws_l.column_dimensions["B"].width = 28
    ws_l.column_dimensions["C"].width = 105
    ws_l["B2"] = "Publications intermédiaires - CGF Gestion"
    ws_l["B2"].font = Font(name="Calibri", bold=True, size=18, color=NAVY)
    ws_l["B3"] = ("Fichier de suivi à charger dans l'application « Fiche de valeur BRVM », au même titre que le classeur "
                  "des états financiers et la base de données SMF.")
    ws_l["B3"].font = Font(name="Calibri", italic=True, size=10, color="555555")
    lines = [
        ("Objet", "Suivre, titre par titre, les indicateurs publiés en cours d'exercice : T1 (1er trimestre), S1 (1er semestre) et T3 (3e trimestre). "
                  "L'application en tire la dynamique de l'exercice en cours : croissance cumulée face à la même période de l'an dernier, tendance "
                  "d'une publication à l'autre, 12 mois glissants. Cette dynamique s'ajoute aux états financiers complets du dernier exercice clos."),
        ("Comment renseigner", "1. Choisissez la feuille « Entreprises » ou « Banques » selon le type du titre (feuille « Titres »).\n"
                               "2. Trouvez la ligne du titre, de l'exercice et de la période, ou ajoutez une ligne sous les dernières.\n"
                               "3. Saisissez les montants de la publication en millions de FCFA, cumulés depuis le 1er janvier.\n"
                               "4. Une publication donne N et N-1 : renseignez aussi la ligne de l'exercice précédent (même période).\n"
                               "5. Indiquez le statut, la date de publication et la source (lien richbourse ou brvm.org).\n"
                               "6. Contrôlez la colonne « Contrôle de cohérence » et suivez l'avancement dans « Suivi »."),
        ("Cumul et périodes", "Toutes les valeurs sont cumulées depuis le 1er janvier : T1 = 3 mois, S1 = 6 mois, T3 = 9 mois. Si une société publie "
                              "un trimestre isolé, additionnez les trimestres. Les encours bancaires (crédits, dépôts) sont ceux de la fin de période."),
        ("Unité et signes", "Millions de FCFA. Les charges se saisissent en négatif (frais de gestion, coût du risque). Les autres soldes sont repris "
                            "tels que publiés. Si le document est en milliers ou en milliards, convertissez en millions."),
        ("Retraitements", "Si le comparatif N-1 d'une publication diffère du chiffre publié l'an dernier (retraitement), retenez le chiffre de la "
                          "publication la plus récente et indiquez-le dans la source."),
        ("Statuts", "Publié : chiffres figurant dans le document publié. Provisoire : communiqué ou chiffres non encore certifiés. "
                    "À vérifier : valeur à contrôler."),
        ("Règles de priorité", "Le fichier prime sur la lecture automatique de richbourse.com, qui comble seulement les périodes ou indicateurs absents. "
                               "Un document PDF chargé dans l'application prime sur la lecture automatique."),
        ("12 mois glissants", "12 mois glissants = exercice N-1 + cumul N - cumul N-1 de la même période. Ils neutralisent la saisonnalité et alimentent "
                              "le PER glissant, la croissance et la tendance. Ils exigent le dernier exercice clos dans les états financiers ou la base SMF."),
        ("Où trouver les chiffres", "brvm.org (avis et communiqués), richbourse.com (rubrique Etats financiers : rapports d'activités T1, S1, T3), sites des sociétés."),
        ("Code couleur", "Texte bleu : saisie. Texte noir : formule, ne pas modifier. Fond gris : colonne calculée."),
    ]
    r = 5
    for lab, txt in lines:
        ws_l.cell(r, 2, lab).font = Font(name="Calibri", bold=True, size=11, color=NAVY)
        ws_l.cell(r, 2).alignment = Alignment(vertical="top")
        ws_l.cell(r, 3, txt).font = Font(name="Calibri", size=10)
        ws_l.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")
        ws_l.row_dimensions[r].height = max(30, 14 * (txt.count("\n") + 1 + len(txt) // 115))
        for j in (2, 3):
            ws_l.cell(r, j).border = Border(bottom=thin)
        r += 1
    ws_l.cell(r + 1, 2, "Société de gestion collective au capital de 500 000 000 FCFA - N° agrément CREPMF SG-003/2001").font = \
        Font(name="Calibri", size=8, color="777777")
    xl._sheets = [xl[n] for n in ["Lisez-moi", "Entreprises", "Banques", "Suivi", "Titres"]]
    xl.active = 0
    buf = io.BytesIO()
    xl.save(buf)
    return buf.getvalue()


def main(argv=None) -> int:
    import argparse
    from pathlib import Path
    ap = argparse.ArgumentParser(description="Génère le classeur des publications intermédiaires (T1, S1, T3)")
    ap.add_argument("--data", help="Fichier JSON des publications collectées (pré-remplissage)")
    ap.add_argument("--out", default="Publications_Intermediaires_CGF.xlsx")
    a = ap.parse_args(argv)
    records = json.loads(Path(a.data).read_text(encoding="utf-8")) if a.data else None
    Path(a.out).write_bytes(build_template(records))
    print(f"Fichier écrit : {a.out}" + (f" ({len(records)} publications)" if records else ""))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
