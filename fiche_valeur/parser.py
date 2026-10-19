"""Lecture du classeur « Etats_Financiers_2021_2025.xlsx ».

Structure attendue (celle du classeur de référence) :

* une feuille ``SYNTHÈSE`` (facultative) et une feuille par société, nommée
  par son ticker BRVM (SGBC, SNTS, PALC...) ;
* en tête de feuille, un titre indiquant l'unité (« en millions FCFA »,
  « en milliers FCFA », « en FCFA »...) ;
* une ligne d'en-tête dont la première cellule vaut ``Poste`` suivie des
  années (2021...2025), puis éventuellement des colonnes de variation ;
* le compte de résultat, puis une section ``BILAN`` (actif / passif), et le
  cas échéant ``HORS BILAN`` et ``RATIOS`` ;
* les charges sont en négatif.

Trois familles de modèles sont reconnues automatiquement : banques (plan
comptable bancaire BCEAO), sociétés commerciales et industrielles
(SYSCOHADA, soldes intermédiaires de gestion) et groupes consolidés IFRS.

Le lecteur est tolérant : les libellés sont reconnus par leur sens (sans
accents, codes SYSCOHADA ou parenthèses), les unités sont harmonisées année
par année lorsqu'une colonne a été saisie dans une autre unité que celle du
titre, et les incohérences internes sont signalées plutôt que masquées.
"""
from __future__ import annotations

import io
import math
import re
from dataclasses import dataclass, field
from typing import Callable

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from . import config
from .utils import is_num, norm_header, norm_label, norm_label_core, strip_accents, to_float

# ---------------------------------------------------------------------------
# Modèle de données
# ---------------------------------------------------------------------------


@dataclass
class Line:
    label: str
    key: str
    core: str
    section: str
    values: dict[int, float]
    row: int


@dataclass
class CompanyData:
    ticker: str
    name: str
    sheet: str
    kind: str                      # "banque" | "syscohada" | "ifrs"
    consolidated: bool
    sector: str
    country: str
    currency: str
    unit_label: str
    multiplier: float
    years: list[int]
    lines: list[Line]
    fields: dict[str, dict[int, float]] = field(default_factory=dict)
    field_labels: dict[str, str] = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)
    estimated: set = field(default_factory=set)

    # Accès pratique ------------------------------------------------------
    def get(self, name: str, year: int | None = None):
        series = self.fields.get(name, {})
        if year is None:
            return series
        return series.get(year)

    def last_year(self, name: str | None = None) -> int | None:
        if name:
            ys = [y for y, v in self.fields.get(name, {}).items() if is_num(v)]
            return max(ys) if ys else None
        return max(self.years) if self.years else None

    @property
    def is_bank(self) -> bool:
        return self.kind == "banque"

    @property
    def revenue_field(self) -> str:
        return "pnb" if self.is_bank else "ca"

    @property
    def net_income_field(self) -> str:
        return "rn_pg" if self.consolidated and self.fields.get("rn_pg") else "rn"

    @property
    def equity_field(self) -> str:
        return "cp_pg" if self.consolidated and self.fields.get("cp_pg") else "capitaux_propres"


@dataclass
class Workbook:
    companies: dict[str, CompanyData]
    synthese_sectors: dict[str, str]
    notes: list[str]
    filename: str = ""

    @property
    def tickers(self) -> list[str]:
        return sorted(self.companies)


# ---------------------------------------------------------------------------
# Règles de reconnaissance des libellés
# ---------------------------------------------------------------------------
# Une règle est un couple (mode, motif) appliqué à la forme normalisée du
# libellé : "eq" égalité, "sw" commence par, "in" contient, "all" contient
# tous les fragments d'un tuple. Les exclusions éliminent un libellé.

Rule = tuple[str, object]


def _r(*rules: Rule) -> list[Rule]:
    return [(m, tuple(norm_label(x) for x in p) if isinstance(p, tuple) else norm_label(p))
            for m, p in rules]


RATIO_WORDS = ["TAUX", "VARIATION", "CROISSANCE", "EVOLUTION", "PAR ACTION", " EN %",
               "POURCENTAGE", "RATIO", "MARGE NETTE"]

FIELD_RULES: dict[str, dict] = {
    # ---- Compte de résultat -------------------------------------------
    "pnb": {"section": "pnl", "rules": _r(("sw", "PRODUIT NET BANCAIRE"), ("eq", "PNB")),
            "exclude": RATIO_WORDS},
    "ca": {"section": "pnl", "rules": _r(
        ("eq", "CHIFFRE D'AFFAIRES"), ("sw", "CHIFFRE D'AFFAIRES"),
        ("sw", "CHIFFRES D'AFFAIRES"), ("sw", "REVENUS DES ACTIVITES ORDINAIRES"),
        ("sw", "PRODUITS DES ACTIVITES ORDINAIRES"), ("eq", "REVENUS"),
        ("sw", "REVENUS CONSOLIDES"), ("sw", "CHIFFRE D AFFAIRES CONSOLIDE")),
        "exclude": RATIO_WORDS + ["EXPORT", "/"]},
    "ventes_marchandises": {"section": "pnl", "rules": _r(("sw", "VENTES DE MARCHANDISES"))},
    "achats_marchandises": {"section": "pnl", "rules": _r(("sw", "ACHATS DE MARCHANDISES"))},
    "marge_commerciale": {"section": "pnl", "rules": _r(("sw", "MARGE COMMERCIALE"), ("sw", "MARGE BRUTE SUR MARCHANDISES"))},
    "production": {"section": "pnl", "rules": _r(("sw", "PRODUCTION DE L'EXERCICE"), ("sw", "PRODUCTION VENDUE"),
                                                  ("sw", "VENTES DE PRODUITS FABRIQUES"))},
    "va": {"section": "pnl", "rules": _r(("sw", "VALEUR AJOUTEE")), "exclude": RATIO_WORDS},
    "charges_personnel": {"section": "pnl", "rules": _r(("sw", "CHARGES DE PERSONNEL"), ("sw", "FRAIS DE PERSONNEL"),
                                                         ("sw", "CHARGES DU PERSONNEL"), ("in", "CHARGES DE PERSONNEL"))},
    "ebe": {"section": "pnl", "rules": _r(("sw", "EXCEDENT BRUT D'EXPLOITATION"), ("eq", "EBITDA"),
                                           ("sw", "EBITDA"), ("eq", "EBE"), ("sw", "EBE ")),
            "exclude": RATIO_WORDS},
    "dotations": {"section": "pnl", "rules": _r(
        ("sw", "DOTATIONS AUX AMORTISSEMENTS"), ("sw", "DOTATION AUX AMORTISSEMENTS"),
        ("sw", "DOTATIONS AUX AMORT"), ("all", ("DOTATIONS", "AMORTISSEMENT")),
        ("sw", "AMORTISSEMENTS ET PROVISIONS"), ("sw", "AMORTISSEMENTS"),
        ("sw", "DEPRECIATIONS ET AMORTISSEMENTS")),
        "exclude": ["REPRISES", "REPRISE"]},
    "interets_produits": {"section": "pnl", "rules": _r(("sw", "INTERETS ET PRODUITS ASSIMILES"),
                                                         ("sw", "PRODUITS D'INTERETS"), ("sw", "INTERETS PERCUS"))},
    "interets_charges": {"section": "pnl", "rules": _r(("sw", "INTERETS ET CHARGES ASSIMILEES"),
                                                        ("sw", "CHARGES D'INTERETS"), ("sw", "INTERETS VERSES"))},
    "marge_interets": {"section": "pnl", "rules": _r(("sw", "MARGE D'INTERET"), ("sw", "MARGE NETTE D'INTERET"),
                                                      ("sw", "MARGE D'INTERETS"))},
    "commissions_produits": {"section": "pnl", "rules": _r(("sw", "COMMISSIONS PRODUITS"), ("sw", "COMMISSIONS PERCUES"),
                                                            ("sw", "PRODUITS DE COMMISSIONS"))},
    "commissions_charges": {"section": "pnl", "rules": _r(("sw", "COMMISSIONS CHARGES"), ("sw", "COMMISSIONS VERSEES"),
                                                           ("sw", "CHARGES DE COMMISSIONS"))},
    "charges_generales": {"section": "pnl", "rules": _r(("sw", "CHARGES GENERALES D'EXPLOITATION"),
                                                         ("sw", "FRAIS GENERAUX"), ("sw", "CHARGES GENERALES"))},
    "rbe": {"section": "pnl", "rules": _r(("sw", "RESULTAT BRUT D'EXPLOITATION"), ("eq", "RBE"))},
    "cout_risque": {"section": "pnl", "rules": _r(("sw", "COUT DU RISQUE"), ("sw", "COUT NET DU RISQUE"))},
    "re": {"section": "pnl", "rules": _r(
        ("eq", "RESULTAT D'EXPLOITATION"), ("sw", "RESULTAT D'EXPLOITATION"),
        ("sw", "RESULTAT OPERATIONNEL"), ("eq", "EBIT"), ("sw", "RESULTAT D EXPLOITATION")),
        "exclude": RATIO_WORDS + ["BRUT"]},
    "rf": {"section": "pnl", "rules": _r(("sw", "RESULTAT FINANCIER"))},
    "frais_financiers": {"section": "pnl", "rules": _r(("sw", "FRAIS FINANCIERS"), ("sw", "CHARGES FINANCIERES"),
                                                        ("sw", "COUT DE L'ENDETTEMENT"))},
    "rao": {"section": "pnl", "rules": _r(("sw", "RESULTAT DES ACTIVITES ORDINAIRES"))},
    "rhao": {"section": "pnl", "rules": _r(("sw", "RESULTAT HORS ACTIVITES ORDINAIRES"),
                                            ("sw", "RESULTAT EXCEPTIONNEL"))},
    "rai": {"section": "pnl", "rules": _r(("sw", "RESULTAT AVANT IMPOT"), ("sw", "RESULTAT COURANT AVANT IMPOT"))},
    "impots": {"section": "pnl", "rules": _r(("sw", "IMPOTS SUR LE RESULTAT"), ("sw", "IMPOT SUR LE RESULTAT"),
                                              ("sw", "IMPOTS SUR LES BENEFICES"), ("sw", "IMPOT SUR LES BENEFICES"),
                                              ("sw", "IMPOT SUR LES SOCIETES"))},
    "rn_pg": {"section": "pnl", "rules": _r(
        ("all", ("RESULTAT", "PART DU GROUPE")), ("all", ("RESULTAT", "ATTRIBUABLE", "MERE")),
        ("all", ("RESULTAT", "ATTRIBUABLE", "PROPRIETAIRES")), ("all", ("RESULTAT", "ATTRIBUABLE", "GROUPE")),
        ("all", ("RESULTAT NET", "ACTIONNAIRES"))),
        "exclude": ["MINORIT", "NE DONNANT PAS", "PAR ACTION", "HORS GROUPE"]},
    "rn": {"section": "pnl", "rules": _r(
        ("eq", "RESULTAT NET"), ("eq", "RESULTAT NET CONSOLIDE"), ("sw", "RESULTAT NET DE L'EXERCICE"),
        ("sw", "RESULTAT NET CONSOLIDE"), ("eq", "RESULTAT DE L'EXERCICE"), ("sw", "BENEFICE OU PERTE"),
        ("sw", "RESULTAT NET")),
        "exclude": RATIO_WORDS + ["PART DU GROUPE", "MINORIT", "ATTRIBUABLE", "BANCAIRE"]},
    # ---- Bilan ----------------------------------------------------------
    "total_actif": {"section": "bilan", "rules": _r(
        ("eq", "TOTAL ACTIF"), ("sw", "TOTAL ACTIF"), ("sw", "TOTAL DE L'ACTIF"), ("sw", "TOTAL GENERAL ACTIF"),
        ("eq", "TOTAL BILAN"), ("sw", "TOTAL BILAN"), ("sw", "TOTAL DU BILAN"), ("eq", "TOTAL GENERAL")),
        "exclude": ["IMMOBILISE", "CIRCULANT", "NON COURANT", "COURANTS", "TRESORERIE"]},
    "total_passif": {"section": "bilan", "rules": _r(("eq", "TOTAL PASSIF"), ("sw", "TOTAL PASSIF"),
                                                      ("sw", "TOTAL DU PASSIF"), ("sw", "TOTAL GENERAL PASSIF")),
                     "exclude": ["CIRCULANT", "COURANT", "TRESORERIE"]},
    "actif_immobilise": {"section": "bilan", "rules": _r(("sw", "ACTIF IMMOBILISE"), ("sw", "TOTAL ACTIF IMMOBILISE"),
                                                          ("sw", "ACTIFS NON COURANTS"), ("sw", "TOTAL ACTIFS NON COURANTS"))},
    "actif_circulant": {"section": "bilan", "rules": _r(("sw", "ACTIF CIRCULANT"), ("sw", "TOTAL ACTIF CIRCULANT"),
                                                         ("sw", "ACTIFS COURANTS"), ("sw", "TOTAL ACTIFS COURANTS"))},
    "stocks": {"section": "bilan", "rules": _r(("sw", "STOCKS"), ("sw", "STOCK"))},
    "creances_clients": {"section": "bilan", "rules": _r(("sw", "CLIENTS"), ("sw", "CREANCES CLIENTS"),
                                                          ("sw", "CLIENTS ET COMPTES RATTACHES")),
                         "exclude": ["AVANCES", "CLIENTELE"]},
    "tresorerie_actif": {"section": "bilan", "rules": _r(("sw", "TRESORERIE ACTIF"), ("sw", "TOTAL TRESORERIE ACTIF"),
                                                          ("sw", "TRESORERIE ET EQUIVALENTS"), ("sw", "DISPONIBILITES"),
                                                          ("sw", "BANQUES CHEQUES POSTAUX CAISSE"))},
    "cp_pg": {"section": "bilan", "rules": _r(
        ("all", ("CAPITAUX PROPRES", "PART DU GROUPE")), ("all", ("CAPITAUX PROPRES", "MERE")),
        ("all", ("CAPITAUX PROPRES", "PROPRIETAIRES")), ("all", ("CAPITAUX PROPRES", "ACTIONNAIRES"))),
        "exclude": ["MINORIT", "NE DONNANT PAS", "TOTAL CAPITAUX PROPRES ET"]},
    "capitaux_propres": {"section": "bilan", "rules": _r(
        ("eq", "CAPITAUX PROPRES"), ("sw", "TOTAL CAPITAUX PROPRES"), ("sw", "CAPITAUX PROPRES ET RESSOURCES ASSIMILEES"),
        ("eq", "CAPITAUX PROPRES TOTAL"), ("sw", "TOTAL DES CAPITAUX PROPRES"), ("eq", "FONDS PROPRES"),
        ("sw", "CAPITAUX PROPRES CONSOLIDES"), ("sw", "CAPITAUX PROPRES")),
        "exclude": ["PART DU GROUPE", "MERE", "MINORIT", "ET DETTES", "ET PASSIF", "PROPRIETAIRES",
                    "RENTABILITE", "VARIATION", "/"]},
    "capital": {"section": "bilan", "rules": _r(("eq", "CAPITAL"), ("sw", "CAPITAL SOCIAL"), ("eq", "CAPITAL APPELE")),
                "exclude": ["NON APPELE", "PROPRES", "PRIME"]},
    "reserves": {"section": "bilan", "rules": _r(("eq", "RESERVES"), ("sw", "RESERVES CONSOLIDEES"),
                                                  ("sw", "RESERVES LIBRES"), ("sw", "RESERVES INDISPONIBLES"))},
    "report_a_nouveau": {"section": "bilan", "rules": _r(("sw", "REPORT A NOUVEAU"))},
    "dettes_financieres": {"section": "bilan", "rules": _r(
        ("sw", "DETTES FINANCIERES"), ("sw", "EMPRUNTS ET DETTES FINANCIERES"), ("sw", "TOTAL DETTES FINANCIERES"),
        ("sw", "EMPRUNTS"), ("sw", "DETTES FINANCIERES NON COURANTES")),
        "exclude": ["CLIENTELE", "INTERBANCAIRE", "ETABLISSEMENTS DE CREDIT", "COURANTES ET"]},
    "passif_circulant": {"section": "bilan", "rules": _r(("sw", "PASSIF CIRCULANT"), ("sw", "TOTAL PASSIF CIRCULANT"),
                                                          ("sw", "PASSIFS COURANTS"), ("sw", "TOTAL PASSIFS COURANTS"))},
    "fournisseurs": {"section": "bilan", "rules": _r(("sw", "FOURNISSEURS"), ("sw", "DETTES FOURNISSEURS"))},
    "tresorerie_passif": {"section": "bilan", "rules": _r(("sw", "TRESORERIE PASSIF"), ("sw", "TOTAL TRESORERIE PASSIF"),
                                                           ("sw", "CONCOURS BANCAIRES"), ("sw", "BANQUES CREDITS DE TRESORERIE"))},
    "credits_clientele": {"section": "bilan", "rules": _r(
        ("sw", "CREANCES SUR LA CLIENTELE"), ("sw", "CREDITS A LA CLIENTELE"), ("sw", "PRETS ET CREANCES SUR LA CLIENTELE"),
        ("sw", "PRETS ET AVANCES A LA CLIENTELE"), ("sw", "CREDITS ET AVANCES A LA CLIENTELE"))},
    "depots_clientele": {"section": "bilan", "rules": _r(
        ("sw", "DETTES A L'EGARD DE LA CLIENTELE"), ("sw", "DEPOTS DE LA CLIENTELE"), ("sw", "DETTES ENVERS LA CLIENTELE"),
        ("sw", "DEPOTS DE CLIENTELE"), ("sw", "DEPOTS DE LA CLIENTELE"), ("sw", "DEPOTS CLIENTELE"))},
    "creances_interbancaires": {"section": "bilan", "rules": _r(
        ("sw", "CREANCES INTERBANCAIRES"), ("sw", "CREANCES SUR LES ETABLISSEMENTS DE CREDIT"))},
    "dettes_interbancaires": {"section": "bilan", "rules": _r(
        ("sw", "DETTES INTERBANCAIRES"), ("sw", "DETTES ENVERS LES ETABLISSEMENTS DE CREDIT"))},
    "titres": {"section": "bilan", "rules": _r(("sw", "OBLIGATIONS ET AUTRES TITRES"), ("sw", "TITRES DE PLACEMENT"),
                                                ("sw", "TITRES D'INVESTISSEMENT"))},
}

# Indicateurs publiés (section RATIOS ou ailleurs) : valeurs relatives
PUBLISHED_RULES: dict[str, list[Rule]] = {
    "roe_pub": _r(("eq", "ROE"), ("sw", "ROE "), ("sw", "RENTABILITE DES FONDS PROPRES"),
                  ("sw", "RENTABILITE DES CAPITAUX PROPRES"), ("sw", "RETURN ON EQUITY")),
    "roa_pub": _r(("eq", "ROA"), ("sw", "ROA "), ("sw", "RENTABILITE DES ACTIFS"), ("sw", "RETURN ON ASSETS")),
    "coef_expl_pub": _r(("sw", "COEFFICIENT D'EXPLOITATION"), ("sw", "COEF D'EXPLOITATION"),
                        ("sw", "COEFFICIENT D EXPLOITATION"), ("in", "COST INCOME")),
    "npl_pub": _r(("in", "CREANCES EN SOUFFRANCE"), ("eq", "NPL"), ("sw", "NPL"), ("in", "CREANCES DOUTEUSES"),
                  ("in", "TAUX DE DEGRADATION")),
    "solvabilite_pub": _r(("in", "SOLVABILITE"), ("in", "RATIO DE FONDS PROPRES")),
    "marge_nette_pub": _r(("sw", "MARGE NETTE"),),
}

# Données par action (pas de multiplicateur d'unité)
PER_SHARE_RULES: dict[str, list[Rule]] = {
    "dps": _r(("in", "DIVIDENDE PAR ACTION"), ("in", "DIVIDENDE NET PAR ACTION"), ("in", "DIVIDENDE BRUT PAR ACTION"),
              ("in", "DIVIDENDES PAR ACTION"), ("eq", "DPA"), ("sw", "DPA "), ("in", "DIVIDENDE UNITAIRE")),
    "bpa_pub": _r(("in", "BENEFICE PAR ACTION"), ("in", "RESULTAT PAR ACTION"), ("eq", "BPA"), ("sw", "BPA ")),
}
COUNT_RULES: dict[str, list[Rule]] = {
    "nb_actions": _r(("in", "NOMBRE D'ACTIONS"), ("in", "NOMBRE DE TITRES"), ("in", "NOMBRE D ACTIONS")),
}
AMOUNT_ANYWHERE_RULES: dict[str, list[Rule]] = {
    "dividendes_distribues": _r(("sw", "DIVIDENDES DISTRIBUES"), ("in", "DIVIDENDES VERSES"),
                                ("sw", "DIVIDENDES PROPOSES"), ("sw", "MONTANT DES DIVIDENDES")),
}


def _match(line: Line, rules: list[Rule], exclude: list[str] | None = None) -> bool:
    forms = (line.key, line.core)
    if exclude:
        excl = [norm_label(e) if e not in ("/", " EN %") else e for e in exclude]
        raw = strip_accents(line.label).upper()
        for e in excl:
            if not e:
                continue
            if e in ("/", " EN %"):
                if e.strip() in raw:
                    return False
            elif any(e in f for f in forms):
                return False
    for mode, pat in rules:
        for f in forms:
            if not f:
                continue
            if mode == "eq" and f == pat:
                return True
            if mode == "sw" and f.startswith(pat):
                return True
            if mode == "in" and pat in f:
                return True
            if mode == "all" and all(p in f for p in pat):
                return True
    return False


def _find(lines: list[Line], rules: list[Rule], exclude=None, sections=None) -> Line | None:
    """Premier libellé correspondant, en respectant la priorité des règles."""
    pool = [ln for ln in lines if (sections is None or ln.section in sections) and ln.values]
    for rule in rules:
        for ln in pool:
            if _match(ln, [rule], exclude):
                return ln
    return None


# ---------------------------------------------------------------------------
# Évaluation minimale des formules (si le classeur n'a pas été recalculé)
# ---------------------------------------------------------------------------
_REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
_SUM = re.compile(r"SUM\(\s*\$?([A-Z]{1,3})\$?(\d+)\s*:\s*\$?([A-Z]{1,3})\$?(\d+)\s*\)", re.I)
_SAFE = re.compile(r"^[0-9eE+\-*/(). ]*$")


class _FormulaSheet:
    def __init__(self, ws_values, ws_formulas):
        self.v = ws_values
        self.f = ws_formulas
        self.cache: dict[str, float | None] = {}

    def value(self, coord: str, depth: int = 0):
        if coord in self.cache:
            return self.cache[coord]
        val = self.v[coord].value
        if is_num(val):
            self.cache[coord] = float(val)
            return float(val)
        num = to_float(val) if isinstance(val, str) else None
        if num is not None:
            self.cache[coord] = num
            return num
        form = self.f[coord].value if self.f is not None else None
        res = None
        if isinstance(form, str) and form.startswith("=") and depth < 25:
            res = self._eval(form[1:], depth + 1)
        self.cache[coord] = res
        return res

    def _eval(self, expr: str, depth: int):
        expr = expr.upper()

        def sum_repl(m):
            c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
            total = 0.0
            for ci in range(column_index_from_string(c1), column_index_from_string(c2) + 1):
                for ri in range(min(r1, r2), max(r1, r2) + 1):
                    v = self.value(f"{get_column_letter(ci)}{ri}", depth)
                    if is_num(v):
                        total += v
            return f"({total!r})"

        expr = _SUM.sub(sum_repl, expr)

        def ref_repl(m):
            v = self.value(f"{m.group(1)}{m.group(2)}", depth)
            return f"({v if is_num(v) else 0.0!r})"

        expr = _REF.sub(ref_repl, expr)
        if not _SAFE.match(expr):
            return None
        try:
            out = eval(expr, {"__builtins__": {}}, {})  # expression arithmétique filtrée
            return float(out) if is_num(out) else None
        except Exception:
            return None


# ---------------------------------------------------------------------------
# Unités
# ---------------------------------------------------------------------------

def _detect_unit(texts: list[str]) -> tuple[float | None, str, str]:
    """Multiplicateur, devise et libellé d'unité lus dans le titre de feuille."""
    for t in texts:
        if not t:
            continue
        u = strip_accents(str(t)).upper()
        currency = "USD" if re.search(r"\bUSD\b|DOLLAR|\$", u) else (
            "EUR" if re.search(r"\bEUR\b|EUROS?\b", u) else "XOF")
        if re.search(r"MILLIARDS?|\bMDS?\b|\bMRDS?\b", u):
            return 1e9, currency, str(t)
        if re.search(r"MILLIONS?|\bMIO\b|\bM ?FCFA\b|\bM ?F\b|\bMUSD\b", u):
            return 1e6, currency, str(t)
        if re.search(r"MILLIERS?|\bK ?FCFA\b|\bKF\b|\bK ?USD\b|\bKUSD\b", u):
            return 1e3, currency, str(t)
        if re.search(r"FCFA|F CFA|\bXOF\b|\bCFA\b|\bUSD\b|\bEUR", u):
            return 1.0, currency, str(t)
    return None, "XOF", ""


# ---------------------------------------------------------------------------
# Lecture d'une feuille société
# ---------------------------------------------------------------------------
YEAR_RE = re.compile(r"^(20\d{2}|19\d{2})(?:\.0+)?$")


def _year_of(value) -> int | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool) and is_num(value):
        v = int(value)
        return v if 1990 <= v <= 2100 and float(value) == v else None
    if isinstance(value, str):
        s = value.strip()
        m = YEAR_RE.match(s)
        if m:
            return int(m.group(1))
        m = re.match(r"^(?:EXERCICE|FY|EX\.?)\s*(20\d{2})$", strip_accents(s).upper())
        if m:
            return int(m.group(1))
        m = re.match(r"^(?:31[/.-]12[/.-])(20\d{2})$", s)
        if m:
            return int(m.group(1))
    if hasattr(value, "year") and hasattr(value, "month"):
        return int(value.year)
    return None


SECTION_HEADERS = [
    ("hors_bilan", ["HORS BILAN", "ENGAGEMENTS"]),
    ("ratios", ["RATIOS", "INDICATEURS", "RATIO "]),
    ("notes", ["NOTES", "SOURCES", "SOURCE ", "COMMENTAIRE", "REMARQUE"]),
    ("bilan", ["BILAN", "ACTIF", "PASSIF", "SITUATION FINANCIERE", "ETAT DE LA SITUATION FINANCIERE"]),
    ("pnl", ["COMPTE DE RESULTAT", "COMPTE DE RESULTATS", "ETAT DU RESULTAT", "SOLDES INTERMEDIAIRES"]),
]


def _section_of_header(key: str, current: str) -> str:
    for sec, words in SECTION_HEADERS:
        for w in words:
            if key == w or key.startswith(w + " ") or key.startswith(w):
                # « ACTIF IMMOBILISE » ou « PASSIF CIRCULANT » sans valeur reste du bilan
                return sec
    return current


def _read_sheet(ws_values, ws_formulas) -> tuple[list[Line], list[int], list[str]]:
    rows = list(ws_values.iter_rows(min_row=1, max_row=min(ws_values.max_row, 400)))
    head_texts: list[str] = []
    candidates: list[tuple[int, dict[int, int], bool]] = []
    for i, row in enumerate(rows[:15]):
        first = row[0].value if row else None
        if isinstance(first, str):
            head_texts.append(first)
        cand: dict[int, int] = {}
        for cell in row[1:]:
            y = _year_of(cell.value)
            if y is not None:
                cand.setdefault(y, cell.column)
        is_poste = isinstance(first, str) and norm_label(first) in {
            "POSTE", "POSTES", "LIBELLE", "LIBELLES", "RUBRIQUE", "RUBRIQUES", "INTITULE", "INTITULES",
            "EN MILLIONS FCFA", "EN MILLIERS FCFA", "EN FCFA"}
        if cand:
            candidates.append((i, cand, is_poste))
    chosen = next((c for c in candidates if c[2]), None) or next(
        (c for c in candidates if len(c[1]) >= 2), None)
    if chosen is None:
        return [], [], head_texts
    header_idx, year_cols, _ = chosen
    fsheet = _FormulaSheet(ws_values, ws_formulas)
    # textes de l'en-tête (unité parfois indiquée en tête de colonne)
    for cell in rows[header_idx]:
        if isinstance(cell.value, str):
            head_texts.append(cell.value)

    lines: list[Line] = []
    section = "pnl"
    for row in rows[header_idx + 1:]:
        if not row:
            continue
        label = row[0].value
        if label is None or (isinstance(label, str) and not label.strip()):
            continue
        if not isinstance(label, str):
            label = str(label)
        key = norm_label(label)
        core = norm_label_core(label)
        vals: dict[int, float] = {}
        for y, col in year_cols.items():
            coord = f"{get_column_letter(col)}{row[0].row}"
            v = fsheet.value(coord)
            if is_num(v):
                vals[y] = float(v)
        if not vals:
            section = _section_of_header(key, section)
            continue
        # Lignes porteuses de valeurs mais qui ouvrent une section (rare)
        if key in ("BILAN", "HORS BILAN", "RATIOS"):
            section = _section_of_header(key, section)
        lines.append(Line(label=label.strip(), key=key, core=core, section=section,
                          values=vals, row=row[0].row))
    years = sorted({y for ln in lines for y in ln.values})
    return lines, years, head_texts


# ---------------------------------------------------------------------------
# Harmonisation des unités année par année
# ---------------------------------------------------------------------------
SCALE_FIELDS = ["ca", "pnb", "total_actif", "capitaux_propres", "va", "charges_personnel",
                "charges_generales", "dettes_financieres", "depots_clientele", "credits_clientele"]
PLAUSIBLE = (2e8, 6e13)  # bornes de vraisemblance en FCFA pour CA / total bilan / fonds propres


def _harmonise(raw: dict[str, dict[int, float]], multiplier: float, fx: float,
               years: list[int]) -> tuple[dict[int, float], float, list[str]]:
    """Facteur correctif par année et multiplicateur global retenu."""
    notes: list[str] = []
    logs: dict[str, dict[int, float]] = {}
    for f in SCALE_FIELDS:
        s = {y: math.log10(abs(v)) for y, v in raw.get(f, {}).items() if is_num(v) and abs(v) > 0}
        if len(s) >= 2:
            logs[f] = s
    year_factor = {y: 1.0 for y in years}
    if logs:
        ref = {f: sorted(s.values())[len(s) // 2] for f, s in logs.items()}
        for y in years:
            devs = [s[y] - ref[f] for f, s in logs.items() if y in s]
            if not devs:
                continue
            devs.sort()
            d = devs[len(devs) // 2]
            if abs(d) >= 2.0:
                k = round(d / 3.0)
                if k != 0 and abs(d - 3 * k) < 0.9:
                    year_factor[y] = 10.0 ** (-3 * k)
    # Choix du multiplicateur global selon la vraisemblance des montants
    best_mult, best_score = multiplier, -1
    for shift in (0, -1, 1, -2, 2, -3, 3):
        mult = multiplier * (1000.0 ** shift)
        score, total = 0, 0
        for f in ("ca", "pnb", "total_actif", "capitaux_propres"):
            for y, v in raw.get(f, {}).items():
                if not is_num(v) or v == 0:
                    continue
                total += 1
                a = abs(v) * year_factor.get(y, 1.0) * mult * fx
                if PLAUSIBLE[0] <= a <= PLAUSIBLE[1]:
                    score += 1
        if total == 0:
            break
        if score > best_score:
            best_mult, best_score = mult, score
    if best_mult != multiplier and best_score > 0:
        notes.append(
            f"Unité du titre de feuille incohérente avec l'ordre de grandeur des montants : "
            f"multiplicateur ajusté de {_fact(multiplier)} à {_fact(best_mult)}.")
    for y, fct in year_factor.items():
        if fct != 1.0:
            notes.append(
                f"Exercice {y} saisi dans une unité différente des autres exercices : "
                f"montants divisés par {_fact(1 / fct)}." if fct < 1 else
                f"Exercice {y} saisi dans une unité différente des autres exercices : "
                f"montants multipliés par {_fact(fct)}.")
    return year_factor, best_mult, notes


def _fact(x: float) -> str:
    return f"{x:,.0f}".replace(",", " ") if x >= 1 else f"{x:g}".replace(".", ",")


# ---------------------------------------------------------------------------
# Lecture du classeur
# ---------------------------------------------------------------------------

def _guess_ticker(sheet_name: str) -> str:
    t = strip_accents(sheet_name).upper().strip()
    t = re.sub(r"[^A-Z0-9]", "", t)
    return t


def _read_synthese(ws) -> tuple[dict[str, str], dict[str, str]]:
    sectors, names = {}, {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 200), values_only=True):
        if not row or not isinstance(row[0], str):
            continue
        t = row[0].strip().upper()
        if re.fullmatch(r"[A-Z]{3,6}", t):
            if len(row) > 1 and isinstance(row[1], str) and row[1].strip():
                sectors.setdefault(t, row[1].strip())
    return sectors, names


def normalise_sector(raw: str | None, ticker: str) -> str:
    ref = config.TICKERS.get(ticker, {}).get("secteur")
    if not raw:
        return ref or "Non classé"
    key = norm_header(raw)
    for s in config.SECTORS:
        if norm_header(s) == key:
            return s
    for alias, target in config.SECTOR_ALIASES.items():
        if norm_header(alias) == key or norm_header(alias) in key:
            return target
    return ref or raw


def _sheet_title_name(texts: list[str]) -> str:
    for t in texts:
        if not t:
            continue
        s = re.split(r"\s[-|:]\s|\(|,|\s-\s", str(t))[0].strip()
        s = re.sub(r"(?i)etats? financiers?|comptes? (sociaux|consolides)|en (milliers|millions|milliards).*", "", s).strip(" -:")
        if len(s) >= 3 and norm_label(s) not in {"POSTE"}:
            return s
    return ""


def parse_workbook(source, filename: str = "", fx_usd_xof: float | None = None) -> Workbook:
    """Lit le classeur des états financiers (chemin, octets ou fichier ouvert)."""
    if isinstance(source, (bytes, bytearray)):
        data = bytes(source)
    elif hasattr(source, "read"):
        data = source.read()
        if hasattr(source, "seek"):
            try:
                source.seek(0)
            except Exception:
                pass
    else:
        with open(source, "rb") as fh:
            data = fh.read()
        filename = filename or str(source)
    wb_v = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    try:
        wb_f = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=False)
    except Exception:
        wb_f = None
    if fx_usd_xof is None:
        fx_usd_xof = config.EUR_XOF / config.EURUSD_FALLBACK

    notes: list[str] = []
    synth_sectors: dict[str, str] = {}
    companies: dict[str, CompanyData] = {}
    for name in wb_v.sheetnames:
        ws = wb_v[name]
        if norm_label(name).startswith("SYNTHESE") or norm_label(name) in {"LISEZ MOI", "README", "SOMMAIRE"}:
            if norm_label(name).startswith("SYNTHESE"):
                synth_sectors, _ = _read_synthese(ws)
            continue
        wsf = wb_f[name] if wb_f is not None and name in wb_f.sheetnames else None
        lines, years, head_texts = _read_sheet(ws, wsf)
        if not lines or not years:
            continue
        ticker = _guess_ticker(name)
        comp = _build_company(ticker, name, lines, years, head_texts, fx_usd_xof)
        if comp is not None:
            companies[ticker] = comp
    for t, comp in companies.items():
        if t in synth_sectors:
            comp.sector = normalise_sector(synth_sectors[t], t)
    if not companies:
        notes.append("Aucune feuille société reconnue : la ligne d'en-tête « Poste | 2021 | 2022... » est introuvable.")
    return Workbook(companies=companies, synthese_sectors=synth_sectors, notes=notes, filename=filename)


def _build_company(ticker: str, sheet: str, lines: list[Line], years: list[int],
                   head_texts: list[str], fx_usd_xof: float) -> CompanyData | None:
    ref = config.TICKERS.get(ticker, {})
    mult, currency, unit_label = _detect_unit(head_texts)
    notes: list[str] = []
    if mult is None:
        mult = 1e6
        unit_label = "unité non précisée (millions FCFA présumés)"
        notes.append("Unité non indiquée dans le titre de la feuille : millions de FCFA présumés, "
                     "puis contrôle de vraisemblance.")
    fx = 1.0
    if currency == "USD":
        fx = fx_usd_xof
        notes.append(f"Comptes publiés en dollars US convertis en FCFA au taux de {fx:,.1f} FCFA/USD."
                     .replace(",", " "))
    elif currency == "EUR":
        fx = config.EUR_XOF

    # 1) Extraction brute des champs standards
    raw: dict[str, dict[int, float]] = {}
    labels: dict[str, str] = {}
    for fname, spec in FIELD_RULES.items():
        sec = spec.get("section")
        hit = _find(lines, spec["rules"], spec.get("exclude"), sections={sec} if sec else None)
        if hit is None:
            # Classeur sans séparation nette des sections : recherche élargie
            hit = _find(lines, spec["rules"], spec.get("exclude"), sections={"pnl", "bilan", "ratios", "notes"})
        if hit is not None:
            raw[fname] = dict(hit.values)
            labels[fname] = hit.label

    # 2) Harmonisation des unités
    year_factor, mult, hnotes = _harmonise(raw, mult, fx, years)
    notes.extend(hnotes)

    def conv(series: dict[int, float]) -> dict[int, float]:
        return {y: v * year_factor.get(y, 1.0) * mult * fx for y, v in series.items() if is_num(v)}

    fields = {k: conv(v) for k, v in raw.items()}

    # 3) Indicateurs publiés, données par action, nombre de titres
    for fname, rules in PUBLISHED_RULES.items():
        hit = _find(lines, rules, None, sections={"ratios", "notes", "hors_bilan", "pnl", "bilan"})
        if hit is not None and fname not in fields:
            vals = {}
            for y, v in hit.values.items():
                vals[y] = v / 100.0 if abs(v) > 1.5 else v
            fields[fname] = vals
            labels[fname] = hit.label
    for fname, rules in PER_SHARE_RULES.items():
        hit = _find(lines, rules, None)
        if hit is not None:
            fields[fname] = {y: v for y, v in hit.values.items() if 0 < abs(v) < 1e6}
            labels[fname] = hit.label
    for fname, rules in COUNT_RULES.items():
        hit = _find(lines, rules, None)
        if hit is not None:
            fields[fname] = {y: v for y, v in hit.values.items() if v > 1000}
            labels[fname] = hit.label
    for fname, rules in AMOUNT_ANYWHERE_RULES.items():
        hit = _find(lines, rules, None)
        if hit is not None:
            fields[fname] = conv(hit.values)
            labels[fname] = hit.label

    # 4) Type de modèle comptable
    if "pnb" in fields:
        kind = "banque"
    elif ticker in config.IFRS_TICKERS or "rn_pg" in fields or "cp_pg" in fields:
        kind = "ifrs"
    else:
        kind = "syscohada"
    consolidated = ticker in config.IFRS_TICKERS or "rn_pg" in fields or "cp_pg" in fields

    name = ref.get("nom") or _sheet_title_name(head_texts) or ticker
    unit_names = {1e9: "milliards", 1e6: "millions", 1e3: "milliers", 1.0: ""}
    cur_name = {"XOF": "FCFA", "USD": "USD", "EUR": "EUR"}.get(currency, currency)
    unit_short = (unit_names.get(mult, f"x{mult:g}") + " " + cur_name).strip()
    comp = CompanyData(
        ticker=ticker, name=name, sheet=sheet, kind=kind, consolidated=consolidated,
        sector=normalise_sector(None, ticker), country=ref.get("pays", "CI"),
        currency=currency, unit_label=unit_short, multiplier=mult * fx,
        years=years, lines=lines, fields=fields, field_labels=labels, notes=notes,
    )
    _derive_and_check(comp)
    return comp


# ---------------------------------------------------------------------------
# Champs dérivés et contrôles de cohérence
# ---------------------------------------------------------------------------

def _derive_and_check(c: CompanyData) -> None:
    f = c.fields
    yrs = c.years

    def setd(name, series, note=None):
        series = {y: v for y, v in series.items() if is_num(v)}
        if series:
            f[name] = series
            c.estimated.add(name)
            if note:
                c.notes.append(note)

    # Marge d'intérêts bancaire
    if c.is_bank and "marge_interets" not in f and "interets_produits" in f and "interets_charges" in f:
        setd("marge_interets", {y: f["interets_produits"][y] + f["interets_charges"][y]
                                for y in yrs if y in f["interets_produits"] and y in f["interets_charges"]})
    # Commissions nettes
    if c.is_bank and "commissions_produits" in f:
        cc = f.get("commissions_charges", {})
        setd("commissions_nettes", {y: f["commissions_produits"][y] + cc.get(y, 0.0)
                                    for y in f["commissions_produits"]})
    # RBE bancaire si absent
    if c.is_bank and "rbe" not in f and "pnb" in f and "charges_generales" in f:
        da = f.get("dotations", {})
        setd("rbe", {y: f["pnb"][y] + f["charges_generales"][y] + da.get(y, 0.0)
                     for y in f["pnb"] if y in f["charges_generales"]},
             "Résultat brut d'exploitation reconstitué (PNB + charges générales + dotations).")
    # Marge commerciale
    if "marge_commerciale" not in f and "ventes_marchandises" in f and "achats_marchandises" in f:
        setd("marge_commerciale", {y: f["ventes_marchandises"][y] + f["achats_marchandises"][y]
                                   for y in f["ventes_marchandises"] if y in f["achats_marchandises"]})
    # EBE estimé : résultat d'exploitation + dotations nettes
    if not c.is_bank and "ebe" not in f and "re" in f and "dotations" in f:
        setd("ebe", {y: f["re"][y] + abs(f["dotations"][y]) for y in f["re"] if y in f["dotations"]},
             "EBE estimé (résultat d'exploitation + dotations aux amortissements), non publié.")
    # Résultat avant impôt
    if "rai" not in f and "rao" in f:
        rh = f.get("rhao", {})
        setd("rai", {y: f["rao"][y] + rh.get(y, 0.0) for y in f["rao"]})
    # Capitaux propres reconstitués
    if "capitaux_propres" not in f and "total_actif" in f and "capital" in f:
        parts = [k for k in ("capital", "reserves", "report_a_nouveau") if k in f]
        rn = f.get("rn", {})
        setd("capitaux_propres", {y: sum(f[k].get(y, 0.0) for k in parts) + rn.get(y, 0.0)
                                  for y in f["capital"]},
             "Capitaux propres reconstitués par somme des postes (capital, réserves, report à nouveau, résultat).")
    if "total_actif" not in f and "total_passif" in f:
        f["total_actif"] = dict(f["total_passif"])

    # Contrôles ---------------------------------------------------------
    rev = c.revenue_field
    for y in list(f.get(rev, {})):
        v = f[rev][y]
        if v <= 0 and not c.is_bank:
            f[rev].pop(y)
            c.notes.append(f"Chiffre d'affaires {y} nul ou négatif dans la source : exclu de l'analyse.")
            continue
        re_ = f.get("re", {}).get(y)
        if not c.is_bank and is_num(re_) and re_ > 0 and v > 0 and re_ > v * 1.05:
            f[rev].pop(y)
            c.notes.append(f"Chiffre d'affaires {y} inférieur au résultat d'exploitation : donnée "
                           f"incohérente exclue (à vérifier sur les états publiés).")
    rn = f.get(c.net_income_field, {})
    for y, v in rn.items():
        ca = f.get(rev, {}).get(y)
        if is_num(ca) and ca > 0 and abs(v) > 1.5 * ca and not c.is_bank:
            c.notes.append(f"Résultat net {y} supérieur à 150% du chiffre d'affaires : élément "
                           f"exceptionnel probable, à vérifier.")
    ta = f.get("total_actif", {})
    tp = f.get("total_passif", {})
    for y in ta:
        if y in tp and ta[y] and abs(ta[y] - tp[y]) / abs(ta[y]) > 0.01:
            c.notes.append(f"Bilan {y} non équilibré (écart actif/passif de "
                           f"{abs(ta[y] - tp[y]) / abs(ta[y]) * 100:.1f}%).".replace(".", ",", 1))
    # Exercices réellement renseignés
    useful = {y for k in ("ca", "pnb", "rn", "rn_pg", "total_actif", "capitaux_propres")
              for y in f.get(k, {})}
    if useful:
        c.years = sorted(useful)
