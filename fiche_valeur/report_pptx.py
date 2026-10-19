"""Rapport PowerPoint (16:9) au format des présentations CGF Gestion.

Le rapport part du modèle maison (assets/templates/modele_cgf_gestion.pptx, issu des présentations CGF
Gestion) : masques COVER, CONTENT, DIVIDER et CLOSE, police du thème, bleu institutionnel, cartes gris-bleu,
bandeau dégradé pour les messages clés. Graphiques natifs modifiables, notes de l'orateur.
"""
from __future__ import annotations

import io
from datetime import date

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from . import charts, config
from .fundamentals import RATIO_META, sig_rows
from .report_pdf import RATIO_ORDER_BANK, RATIO_ORDER_CORP, _fmt_ratio, _var_txt, METHODO
from .utils import (clean_text, fmt_date_fr, fmt_fcfa, fmt_mds, fmt_num, fmt_pct, fmt_price, fmt_pts, fmt_x,
                    is_num)

W, H = Inches(13.333), Inches(7.5)
HEAD = None          # police du thème du modèle (titres)
BODY = None          # police du thème du modèle (texte)


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#").upper())


NAVY, GOLD, INK, MUTED = rgb(config.NAVY), rgb(config.GOLD), rgb(config.INK), rgb(config.MUTED)
PAPER, WHITE, GRIDC = rgb(config.PAPER), rgb("#FFFFFF"), rgb(config.GRID)
SKY, INFO, LIGHTLINE, FOOT = rgb(config.SKY), rgb(config.INFO), rgb(config.GOLD_LIGHT), rgb("#A6A6A6")
DARK, NUMBLUE = rgb(config.NAVY_DARK), rgb("#2A6E9E")
ICONS = config.BRAND_DIR / "icons"
S1, S2, S3, S4, GRAY = rgb(charts.S1), rgb(charts.S2), rgb(charts.S3), rgb(charts.S4), rgb(charts.GRAY)


# ---------------------------------------------------------------------------
# Primitives
# ---------------------------------------------------------------------------

def _rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, name=None):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    if name:
        s.name = name
    return s


def _text(slide, x, y, w, h, text, size=12, bold=False, color=INK, font=BODY, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, italic=False, name=None, line_spacing=1.05):
    tb = slide.shapes.add_textbox(x, y, w, h)
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    paras = text if isinstance(text, list) else [text]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        runs = para if isinstance(para, tuple) else ((para, {}),)
        if isinstance(para, str):
            runs = ((para, {}),)
        for rtext, ropt in runs:
            r = p.add_run()
            r.text = clean_text(rtext)
            f = r.font
            if ropt.get("font", font):
                f.name = ropt.get("font", font)
            f.size = Pt(ropt.get("size", size))
            f.bold = ropt.get("bold", bold)
            f.italic = ropt.get("italic", italic)
            f.color.rgb = ropt.get("color", color)
    return tb


def _bullets(slide, x, y, w, h, items, size=11, color=INK, bullet_color=GOLD, space_after=4):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        r1 = p.add_run()
        r1.text = "●  "
        r1.font.size = Pt(size - 3)
        r1.font.color.rgb = bullet_color
        r2 = p.add_run()
        r2.text = clean_text(it)
        r2.font.size = Pt(size)
        r2.font.color.rgb = color
    return tb


def _table(slide, x, y, w, rows, col_w=None, font=10, header=True, row_h=0.26, bold_first_col=False,
           align_right_from=1, left_cols=()):
    nr, nc = len(rows), max(len(r) for r in rows)
    shape = slide.shapes.add_table(nr, nc, x, y, w, Inches(row_h * nr))
    tbl = shape.table
    tbl.first_row = header
    tbl.horz_banding = False
    if col_w:
        tot = sum(col_w)
        for j, cw in enumerate(col_w):
            tbl.columns[j].width = Emu(int(w * cw / tot))
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j in range(nc):
            val = row[j] if j < len(row) else ""
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.015)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if header and i == 0:
                cell.fill.fore_color.rgb = NAVY
            else:
                cell.fill.fore_color.rgb = PAPER if (i % 2 == (0 if header else 1)) else WHITE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if (j >= align_right_from and j not in left_cols) else PP_ALIGN.LEFT
            r = p.add_run()
            r.text = clean_text(str(val))
            r.font.size = Pt(font)
            r.font.bold = (header and i == 0) or (bold_first_col and j == 0)
            r.font.color.rgb = WHITE if (header and i == 0) else INK
    return shape


def _style_chart(chart, legend=True, font=9):
    chart.font.size = Pt(font)
    chart.font.color.rgb = MUTED
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(font)
    try:
        va = chart.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = GRIDC
        va.major_gridlines.format.line.width = Pt(0.5)
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(font)
        ca = chart.category_axis
        ca.format.line.color.rgb = GRIDC
        ca.tick_labels.font.size = Pt(font)
        ca.has_major_gridlines = False
    except Exception:
        pass


def _bar_chart(slide, x, y, w, h, categories, series, colors, fmt='#,##0.0', title=None, horizontal=False,
               legend=True, labels_last=True, gap=60, point_colors=None):
    cd = CategoryChartData()
    cd.categories = [str(c) for c in categories]
    for name, vals in series:
        cd.add_series(name, [v if is_num(v) else None for v in vals], number_format=fmt)
    ct = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(ct, x, y, w, h, cd)
    ch = gf.chart
    _style_chart(ch, legend=legend and len(series) > 1)
    if title:
        ch.has_title = True
        ch.chart_title.text_frame.text = clean_text(title)
        tp = ch.chart_title.text_frame.paragraphs[0]
        tp.runs[0].font.size = Pt(11)
        tp.runs[0].font.bold = True
        tp.runs[0].font.color.rgb = NAVY
    else:
        ch.has_title = False
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.overlap = -10 if len(series) > 1 else 0
    for i, s in enumerate(plot.series):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = colors[i % len(colors)]
        s.format.line.fill.background()
        s.invert_if_negative = False
    if point_colors:
        for idx, col in point_colors.items():
            pt = plot.series[0].points[idx]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = col
    if labels_last:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = fmt
        dl.number_format_is_linked = False
        dl.font.size = Pt(8)
        dl.font.color.rgb = INK
    ch.value_axis.tick_labels.number_format = fmt.split(";")[0]
    ch.value_axis.tick_labels.number_format_is_linked = False
    if horizontal:
        ch.category_axis.reverse_order = True
        ch.value_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    return ch


def _line_chart(slide, x, y, w, h, categories, series, colors, fmt='#,##0', title=None, widths=None, legend=True,
                date_axis=False):
    cd = CategoryChartData()
    cd.categories = categories
    for name, vals in series:
        cd.add_series(name, [v if is_num(v) else None for v in vals], number_format=fmt)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE, x, y, w, h, cd)
    ch = gf.chart
    _style_chart(ch, legend=legend and len(series) > 1)
    ch.has_title = bool(title)
    if title:
        ch.chart_title.text_frame.text = clean_text(title)
        r = ch.chart_title.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.bold, r.font.color.rgb = Pt(11), True, NAVY
    for i, s in enumerate(ch.plots[0].series):
        s.format.line.color.rgb = colors[i % len(colors)]
        s.format.line.width = Pt((widths or [2.0] * len(series))[i])
        s.smooth = False
        s.marker.style = None
        from pptx.enum.chart import XL_MARKER_STYLE
        s.marker.style = XL_MARKER_STYLE.NONE if date_axis else XL_MARKER_STYLE.CIRCLE
        if not date_axis:
            s.marker.size = 6
            s.marker.format.fill.solid()
            s.marker.format.fill.fore_color.rgb = colors[i % len(colors)]
            s.marker.format.line.color.rgb = WHITE
    ch.value_axis.tick_labels.number_format = fmt
    ch.value_axis.tick_labels.number_format_is_linked = False
    if date_axis:
        ca = ch.category_axis
        ca.tick_labels.number_format = 'mmm yy'
        ca.tick_labels.number_format_is_linked = False
    return ch


def _image(slide, png, x, y, w=None, h=None):
    if not png:
        return None
    return slide.shapes.add_picture(io.BytesIO(png), x, y, width=w, height=h)


# ---------------------------------------------------------------------------
# Cadre des diapositives
# ---------------------------------------------------------------------------

def _picture(slide, path, x, y, w=None, h=None):
    try:
        return slide.shapes.add_picture(str(path), x, y, width=w, height=h)
    except Exception:
        return None


def _icon_circle(slide, icon, x, y, d, fill=None):
    """Pastille ronde bleue et pictogramme blanc, comme dans les présentations CGF Gestion."""
    c = _rect(slide, x, y, d, d, fill or NAVY, MSO_SHAPE.OVAL)
    pad = int(d * 0.25)
    _picture(slide, ICONS / f"{icon}.png", x + pad, y + pad, d - 2 * pad, d - 2 * pad)
    return c


def _card(slide, x, y, w, h, fill=PAPER, line=GRIDC, radius=0.04):
    r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    r.adjustments[0] = radius
    r.fill.solid()
    r.fill.fore_color.rgb = fill
    r.line.color.rgb = line
    r.line.width = Pt(0.75)
    r.shadow.inherit = False
    return r


def _band(slide, x, y, w, h, text, size=12.5):
    """Bandeau dégradé sombre portant le message clé (texte blanc)."""
    _picture(slide, config.BRAND_DIR / "bandeau.jpg", x, y, w, h)
    _text(slide, x + Inches(0.3), y, w - Inches(0.6), h, text, size=size, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)


def _info(slide, x, y, w, h, text, icon="idee", size=12):
    """Encadré bleu clair avec pictogramme (repère pédagogique du modèle)."""
    _card(slide, x, y, w, h, fill=INFO, line=LIGHTLINE)
    d = min(h - Inches(0.3), Inches(0.62))
    _icon_circle(slide, icon, x + Inches(0.3), y + (h - d) // 2, d, fill=GOLD)
    _text(slide, x + Inches(0.3) + d + Inches(0.25), y, w - d - Inches(0.8), h, text, size=size, color=INK,
          anchor=MSO_ANCHOR.MIDDLE)


class Deck:
    """Diapositives construites sur le modèle CGF Gestion."""

    def __init__(self, a):
        self.a = a
        self.prs = Presentation(str(config.PPTX_TEMPLATE)) if config.PPTX_TEMPLATE.exists() else Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.layouts = {l.name: l for l in self.prs.slide_layouts}
        self.blank = self.layouts.get("Vide", self.prs.slide_layouts[6 if len(self.prs.slide_layouts) > 6 else 0])
        self.n = 0
        foot = clean_text(f"FICHE DE VALEUR {a.company.ticker}  |  {a.settings.maison}  |  "
                          f"{fmt_date_fr(a.as_of)}  |  Document interne")
        for name in ("CONTENT", "CONTENT_PANEL"):
            lay = self.layouts.get(name)
            if lay is None:
                continue
            for sh in lay.shapes:
                if sh.has_text_frame and "{PIED_DE_PAGE}" in sh.text_frame.text:
                    sh.text_frame.paragraphs[0].runs[0].text = foot

    def _add(self, layout_name):
        lay = self.layouts.get(layout_name, self.blank)
        s = self.prs.slides.add_slide(lay)
        self.n += 1
        return s

    def _number(self, s):
        _text(s, Inches(12.3), Inches(7.03), Inches(0.7), Inches(0.3), str(self.n), size=9, color=FOOT,
              align=PP_ALIGN.RIGHT, name="Numéro")

    def slide(self, title, subtitle=None, notes=None):
        a = self.a
        s = self._add("CONTENT")
        if s.shapes.title is not None:
            s.shapes.title.text = clean_text(title)
            s.shapes.title.width = Inches(10.4)
        else:
            _text(s, Inches(0.36), Inches(0.22), Inches(10.4), Inches(0.52), title, size=26, bold=True, color=NAVY)
        if subtitle:
            _text(s, Inches(0.36), Inches(0.82), Inches(10.4), Inches(0.32), subtitle, size=12, color=MUTED,
                  name="Sous-titre")
        tg = _rect(s, Inches(11.1), Inches(0.27), Inches(1.87), Inches(0.42),
                   rgb(config.SIGNAL_COLORS.get(a.decision.signal, config.NAVY)), MSO_SHAPE.ROUNDED_RECTANGLE,
                   name="Etiquette signal")
        tf = tg.text_frame
        tf.text = clean_text(f"{a.company.ticker}  |  {a.decision.signal}")
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.size, p.runs[0].font.bold, p.runs[0].font.color.rgb = Pt(11), True, WHITE
        for ph in list(s.placeholders):
            if ph.placeholder_format.type is not None and "SLIDE_NUMBER" in str(ph.placeholder_format.type):
                ph._element.getparent().remove(ph._element)
        self._number(s)
        if notes:
            s.notes_slide.notes_text_frame.text = clean_text(notes)
        return s

    def cover(self):
        a, c, dec, v = self.a, self.a.company, self.a.decision, self.a.valuation
        country = config.COUNTRIES.get(c.country, {}).get("nom", c.country)
        s = self._add("COVER")
        _text(s, Inches(3.25), Inches(1.55), Inches(9.6), Inches(1.4), c.name, size=38, bold=True, color=WHITE,
              anchor=MSO_ANCHOR.BOTTOM)
        _text(s, Inches(3.25), Inches(3.05), Inches(9.3), Inches(0.45), "Fiche de valeur et recommandation", size=22,
              color=SKY)
        _text(s, Inches(3.25), Inches(3.65), Inches(9.3), Inches(0.9),
              [f"{c.ticker}  |  {c.sector}  |  {country}  |  BRVM",
               ((f"Données arrêtées au {fmt_date_fr(a.as_of)}", {"size": 13, "color": rgb("#C9D6E2")}),)],
              size=15, color=WHITE)
        pill = _rect(s, Inches(3.25), Inches(4.75), Inches(6.4), Inches(0.62),
                     rgb(config.SIGNAL_COLORS.get(dec.signal, config.NAVY)), MSO_SHAPE.ROUNDED_RECTANGLE)
        pill.adjustments[0] = 0.5
        tf = pill.text_frame
        tf.text = clean_text(f"{dec.signal}  |  conviction {dec.conviction.lower()}  |  potentiel "
                             f"{fmt_pct(dec.upside, sign=True)}")
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.size, p.runs[0].font.bold, p.runs[0].font.color.rgb = Pt(13), True, WHITE
        _picture(s, config.BRAND_DIR / "logo_cgf.png", Inches(0.42), Inches(6.62), Inches(2.0), Inches(0.38))
        _text(s, Inches(8.4), Inches(6.9), Inches(4.5), Inches(0.3),
              f"{a.settings.maison}  |  Société de gestion agréée par l'AMF-UMOA", size=9, color=rgb("#C9D6E2"),
              align=PP_ALIGN.RIGHT)
        if a.settings.analyste:
            _text(s, Inches(3.25), Inches(5.55), Inches(9), Inches(0.35), f"Analyste : {a.settings.analyste}", size=12,
                  color=rgb("#C9D6E2"))
        s.notes_slide.notes_text_frame.text = clean_text(" ".join(a.texts.get("these", [])))
        return s

    def divider(self, num, title, subtitle, icon="livre"):
        s = self._add("DIVIDER")
        _text(s, Inches(0.9), Inches(2.25), Inches(2.0), Inches(1.5), f"{num:02d}", size=80, bold=True, color=NUMBLUE,
              align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        ln = s.shapes.add_connector(1, Inches(3.05), Inches(2.5), Inches(3.05), Inches(3.5))
        ln.line.color.rgb = GOLD
        ln.line.width = Pt(2)
        _text(s, Inches(3.3), Inches(2.4), Inches(8.6), Inches(0.72), title, size=30, bold=True, color=WHITE)
        _text(s, Inches(3.33), Inches(3.12), Inches(8.6), Inches(0.45), subtitle, size=15, color=SKY)
        _picture(s, ICONS / f"{icon}.png", Inches(11.55), Inches(0.55), Inches(0.72), Inches(0.72))
        return s

    def close(self):
        a = self.a
        s = self._add("CLOSE")
        _text(s, Inches(1.0), Inches(1.7), Inches(7.0), Inches(0.9), "Merci de votre attention", size=40, bold=True,
              color=WHITE)
        _text(s, Inches(1.0), Inches(4.05), Inches(6.6), Inches(0.3), "POUR ALLER PLUS LOIN", size=11, bold=True,
              color=SKY)
        steps = [("calendrier", f"Fiche établie le {fmt_date_fr(a.as_of)} à partir des données disponibles à cette date"),
                 ("loupe", "Détail des calculs dans l'export Excel joint à la fiche"),
                 ("bouclier", "Signal à valider en comité d'investissement avant toute décision")]
        for i, (ic, txt) in enumerate(steps):
            yy = Inches(4.48 + 0.52 * i)
            _icon_circle(s, ic, Inches(1.0), yy, Inches(0.36), fill=NUMBLUE)
            _text(s, Inches(1.5), yy - Inches(0.04), Inches(6.4), Inches(0.44), txt, size=12, color=WHITE,
                  anchor=MSO_ANCHOR.MIDDLE)
        for i, (ic, main, sub) in enumerate(config.CONTACTS):
            yy = Inches(1.72 + 1.08 * i)
            _icon_circle(s, ic, Inches(8.3), yy, Inches(0.56), fill=NUMBLUE)
            _text(s, Inches(9.02), yy - Inches(0.02), Inches(3.9), Inches(0.66),
                  [((main, {"bold": True, "size": 13}),), ((sub, {"size": 10, "color": rgb("#C9D6E2")}),)], color=WHITE)
        _text(s, Inches(1.0), Inches(6.1), Inches(11.3), Inches(0.56), config.DISCLAIMER, size=8.5, italic=True,
              color=rgb("#C9D6E2"))
        _picture(s, config.BRAND_DIR / "logo_cgf_blanc.png", Inches(1.0), Inches(6.72), Inches(1.75), Inches(0.33))
        return s


# ---------------------------------------------------------------------------
# Construction
# ---------------------------------------------------------------------------

def build_pptx(a) -> bytes:
    d = Deck(a)
    c, v, fd, t, m, dec, tx = a.company, a.valuation, a.fundamentals, a.technical, a.macro, a.decision, a.texts
    y = fd.year
    sigcol = rgb(config.SIGNAL_COLORS.get(dec.signal, config.NAVY))
    country = config.COUNTRIES.get(c.country, {}).get("nom", c.country)

    # 1. Couverture ---------------------------------------------------------
    d.cover()

    # 2. Synthèse -------------------------------------------------------------
    s = d.slide("Synthèse et recommandation", f"Données arrêtées au {fmt_date_fr(a.as_of)}",
                notes=" ".join(tx.get("justification", [])))
    kpis = [("tendance", fmt_price(v.price), f"Cours au {a.market.price_date:%d/%m/%Y}" if a.market.price_date else "Cours"),
            ("cible", fmt_price(dec.target), "Juste valeur, cible à 12 mois"),
            ("pourcentage", fmt_pct(dec.upside, sign=True), "Potentiel vers la juste valeur"),
            ("pieces", fmt_pct(v.div_yield), "Rendement net du dividende")]
    cw, gap = Inches(2.99), Inches(0.26)
    for i, (ic, val, lab) in enumerate(kpis):
        x = Inches(0.36) + i * (cw + gap)
        _card(s, x, Inches(1.32), cw, Inches(2.05))
        _icon_circle(s, ic, x + Inches(0.3), Inches(1.55), Inches(0.56))
        _text(s, x + Inches(0.3), Inches(2.2), cw - Inches(0.5), Inches(0.55), val, size=24, bold=True, color=NAVY)
        _text(s, x + Inches(0.3), Inches(2.78), cw - Inches(0.5), Inches(0.5), lab, size=11, color=MUTED)
    these = tx.get("these", [])
    _text(s, Inches(0.36), Inches(3.55), Inches(7.6), Inches(0.32), "THÈSE D'INVESTISSEMENT", size=11, bold=True,
          color=NAVY)
    _text(s, Inches(0.36), Inches(3.9), Inches(7.6), Inches(2.0), these[-1:] if these else [], size=11,
          line_spacing=1.08)
    _text(s, Inches(8.3), Inches(3.55), Inches(4.7), Inches(0.32), "POINTS FORTS", size=11, bold=True, color=NAVY)
    _bullets(s, Inches(8.3), Inches(3.88), Inches(4.7), Inches(1.0), (tx.get("forces") or ["Aucun point saillant."])[:3],
             size=10)
    _text(s, Inches(8.3), Inches(4.9), Inches(4.7), Inches(0.32), "POINTS DE VIGILANCE", size=11, bold=True, color=NAVY)
    _bullets(s, Inches(8.3), Inches(5.22), Inches(4.7), Inches(0.75),
             (tx.get("vigilances") or ["Aucun point saillant."])[:2], size=10, bullet_color=rgb(config.NEG))
    msg = (f"Signal {dec.signal}, conviction {dec.conviction.lower()}"
           + (f", note globale {fmt_num(dec.composite)}/100" if is_num(dec.composite) else "")
           + ". " + clean_text(config.SIGNAL_DEFINITIONS.get(dec.signal, "")))
    _band(s, Inches(0.36), Inches(6.05), Inches(12.61), Inches(0.85), msg, size=11.5)

    d.divider(1, "Analyse boursière et financière", "Profil boursier, soldes intermédiaires de gestion, ratios et pairs",
              icon="institution")

    # 3. Profil boursier --------------------------------------------------------
    s = d.slide("Profil boursier", "Données de marché et évolution du cours", notes=tx.get("technique"))
    liq = t.liquidity if t.available else {}
    rows = [["Indicateur", "Valeur"],
            ["Cours", fmt_price(v.price)], ["Capitalisation", fmt_fcfa(a.market.market_cap)],
            ["Nombre de titres", fmt_num(a.market.shares)], ["Flottant", fmt_pct(a.market.float_pct, 0)],
            ["PER" + (f" {y}" if y else ""), fmt_x(v.pe)], ["P/B", fmt_x(v.pb, 2)], ["BPA" + (f" {y}" if y else ""), fmt_price(v.eps)],
            ["DPA net", fmt_price(v.dps)], ["Rendement net", fmt_pct(v.div_yield)],
            ["Bêta (BRVM-C)", fmt_num(v.beta, 2)],
            ["Perf. depuis le 1er janv.", fmt_pct(t.performance.get("ytd"), sign=True) if t.available else "n.d."],
            ["Perf. 12 mois", fmt_pct(t.performance.get("12m"), sign=True) if t.available else "n.d."],
            ["Volume moyen 3 mois", fmt_num(liq.get("volume_moyen_3m")) if liq else "n.d."],
            ["Séances actives 3 mois", fmt_pct(liq.get("taux_seances_actives"), 0) if liq else "n.d."]]
    _table(s, Inches(0.6), Inches(1.5), Inches(3.9), rows, [1.6, 1.2], font=10.5, row_h=0.33)
    if t.available:
        _price_native(s, t, Inches(4.8), Inches(1.4), Inches(8.0), Inches(5.5), months=18)
    else:
        _text(s, Inches(4.8), Inches(3), Inches(7), Inches(1), "Historique de cours indisponible.", size=14, color=MUTED)

    if c.level in ("complet", "partiel"):
        # 4. Soldes intermédiaires de gestion ---------------------------------------
        years = c.years[-5:]
        s = d.slide("Soldes intermédiaires de gestion", f"En milliards de FCFA, exercices {years[0]}-{years[-1]}",
                    notes=tx.get("activite"))
        rows = [["Mds FCFA"] + [str(yy) for yy in years] + ["Var."]]
        for key, label, ser in sig_rows(c):
            rows.append([label] + [fmt_mds(ser.get(yy)) for yy in years] + [_var_txt(ser, years)])
        rh = min(0.34, 4.9 / max(len(rows), 1))
        _table(s, Inches(0.6), Inches(1.5), Inches(7.4), rows, [2.6] + [1] * (len(years) + 1), font=10, row_h=rh)
        rev = c.fields.get(c.revenue_field, {})
        rn = c.fields.get(c.net_income_field, {})
        _bar_chart(s, Inches(8.25), Inches(1.4), Inches(4.6), Inches(3.6), years,
                   [("PNB" if c.is_bank else "Chiffre d'affaires", [rev.get(yy, 0) / 1e9 if is_num(rev.get(yy)) else None for yy in years]),
                    ("Résultat net", [rn.get(yy) / 1e9 if is_num(rn.get(yy)) else None for yy in years])],
                   [S1, S2], fmt='#,##0', title="Revenus et résultat net (Mds FCFA)", labels_last=False)
        _text(s, Inches(8.25), Inches(5.15), Inches(4.6), Inches(1.8), tx.get("activite", ""), size=10, color=INK)

        # 5. Ratios ------------------------------------------------------------------
        s = d.slide("Rentabilité et structure financière", "Ratios calculés à partir des états financiers publiés",
                    notes="Lecture des ratios : " + tx.get("activite", ""))
        order = RATIO_ORDER_BANK if c.is_bank else RATIO_ORDER_CORP
        rows = [["Ratio"] + [str(yy) for yy in years]]
        for k in order:
            ser = fd.ratios.get(k, {})
            if any(is_num(ser.get(yy)) for yy in years):
                rows.append([RATIO_META[k][0]] + [_fmt_ratio(k, ser.get(yy)) for yy in years])
        rows = rows[:15]
        _table(s, Inches(0.6), Inches(1.5), Inches(7.4), rows, [3.0] + [1] * len(years), font=9.5,
               row_h=min(0.33, 5.3 / len(rows)))
        lines = [("roe", "ROE", S1), ("marge_nette", "Marge nette", S2),
                 ("coef_exploitation", "Coefficient d'exploitation", S3) if c.is_bank else ("marge_ebe", "Marge d'EBE", S3)]
        series = [(lab, [fd.ratios.get(k, {}).get(yy) for yy in years]) for k, lab, _ in lines]
        _line_chart(s, Inches(8.25), Inches(1.4), Inches(4.6), Inches(3.9), [str(yy) for yy in years], series,
                    [col for *_, col in lines], fmt='0%', title="Rentabilité")
        sol = fd.solidity
        _text(s, Inches(8.25), Inches(5.45), Inches(4.6), Inches(1.4),
              [(("Score de solidité ", {"bold": True, "color": NAVY}), (f"{sol.score}/{sol.maximum}", {"bold": True})),
               f"Z'' d'Altman : {fmt_num(sol.altman, 2)} (zone de {sol.altman_zone})" if sol.altman is not None else
               "Tests : rentabilité, efficacité, risque, transformation, solvabilité."], size=11)

        # 6. Solidité et pairs ---------------------------------------------------------
        s = d.slide("Solidité et positionnement sectoriel", f"Comparaison avec les pairs : {fd.peer_label}")
        rows = [["Test de solidité", "Résultat"]] + [[lab, "Validé" if ok else ("Non validé" if ok is False else "n.d.")]
                                                   for lab, ok in sol.tests]
        _table(s, Inches(0.6), Inches(1.5), Inches(5.6), rows, [3.6, 1.2], font=10, row_h=0.34)
        peers = [(p.ticker, p.roe) for p in fd.peers if is_num(p.roe) and -0.5 < p.roe < 1]
        me = fd.ratios.get("roe", {}).get(y)
        if is_num(me):
            peers.append((c.ticker, me))
        peers.sort(key=lambda z: z[1], reverse=True)
        if len(peers) >= 2:
            idx = [i for i, (tk, _) in enumerate(peers) if tk == c.ticker]
            pc = {i: S1 for i in idx}
            _bar_chart(s, Inches(6.6), Inches(1.4), Inches(6.2), Inches(5.4), [tk for tk, _ in peers],
                       [("ROE", [r for _, r in peers])], [GRAY], fmt='0.0%', title=f"ROE {y} : {c.ticker} et ses pairs",
                       horizontal=True, legend=False, gap=45, point_colors=pc)

    else:
        years = c.years[-5:]
        if c.level == "chiffres_cles" and years:
            s = d.slide("Chiffres clés", f"Exercices {years[0]}-{years[-1]}, source : {clean_text(c.data_source)}",
                        notes=tx.get("activite"))
            rev, rn = c.fields.get(c.revenue_field, {}), c.fields.get(c.net_income_field, {})
            eps, dps = c.fields.get("bpa_pub", {}), c.fields.get("dps", {})
            cp = c.fields.get(c.equity_field, {})
            rows = [["Mds FCFA sauf mention"] + [str(yy) for yy in years]]
            for lab, ser, fn in (("PNB" if c.is_bank else "Chiffre d'affaires", rev, fmt_mds), ("Résultat net", rn, fmt_mds),
                                 ("Capitaux propres", cp, fmt_mds), ("BNPA (FCFA)", eps, fmt_num),
                                 ("Dividende net (FCFA)", dps, fmt_num)):
                if any(is_num(ser.get(yy)) for yy in years):
                    rows.append([lab] + [fn(ser.get(yy)) if is_num(ser.get(yy)) else "n.d." for yy in years])
            mn = [(rn[yy] / rev[yy]) if is_num(rn.get(yy)) and is_num(rev.get(yy)) and rev[yy] else None for yy in years]
            if any(is_num(m_) for m_ in mn):
                rows.append(["Marge nette"] + [fmt_pct(m_) if is_num(m_) else "n.d." for m_ in mn])
            _table(s, Inches(0.6), Inches(1.5), Inches(7.4), rows, [2.6] + [1] * len(years), font=10.5, row_h=0.4)
            _bar_chart(s, Inches(8.25), Inches(1.4), Inches(4.6), Inches(3.6), years,
                       [("PNB" if c.is_bank else "Chiffre d'affaires",
                         [rev.get(yy, 0) / 1e9 if is_num(rev.get(yy)) else None for yy in years]),
                        ("Résultat net", [rn.get(yy) / 1e9 if is_num(rn.get(yy)) else None for yy in years])],
                       [S1, S2], fmt='#,##0', title="Revenus et résultat net (Mds FCFA)", labels_last=False)
            _text(s, Inches(8.25), Inches(5.15), Inches(4.6), Inches(1.8), tx.get("activite", ""), size=10, color=INK)
        cov = a.coverage
        s = d.slide("Données disponibles et compléments",
                    f"Couverture des données : {fmt_num(cov.score, 0)}%")
        lab = {"ok": "Disponible", "partiel": "Partiel", "absent": "Absent"}
        rows = [["Donnée", "Statut", "Détail"]] + [[i.name, lab[i.status], clean_text(i.detail)[:70]] for i in cov.items]
        _table(s, Inches(0.6), Inches(1.5), Inches(7.6), rows, [2.3, 1.0, 3.6], font=9.5, row_h=0.42,
               left_cols=(0, 1, 2))
        todo = [f"{i.name} : {clean_text(i.action)}" for i in cov.missing][:6]
        _text(s, Inches(8.5), Inches(1.5), Inches(4.3), Inches(0.35), "Pour compléter l'analyse", size=14, bold=True,
              color=NAVY)
        _bullets(s, Inches(8.5), Inches(1.95), Inches(4.3), Inches(4.9),
                 todo or ["Toutes les données principales sont disponibles."], size=10.5)

    d.divider(2, "Valorisation et marché", "Juste valeur, analyse technique, environnement et actualités",
              icon="balance")

    # 7. Valorisation ----------------------------------------------------------------
    s = d.slide("Valorisation", f"Juste valeur {fmt_price(v.fair_value)}, potentiel {fmt_pct(v.upside, sign=True)}",
                notes=tx.get("valorisation"))
    tw = sum(mm.weight for mm in v.methods) or 1
    rows = [["Méthode", "FCFA / action", "Poids"]] + [[mm.name, fmt_num(mm.value), fmt_pct(mm.weight / tw, 0)]
                                                       for mm in v.methods]
    rows.append(["Juste valeur pondérée", fmt_num(v.fair_value), "100%"])
    _table(s, Inches(0.6), Inches(1.5), Inches(5.4), rows, [2.6, 1.3, 0.8], font=10.5, row_h=0.36)
    hyp = [f"Coût des fonds propres {fmt_pct(v.cost_of_equity)} (taux sans risque "
           f"{fmt_pts(a.settings.taux_sans_risque, 2, sign=False)}, prime {fmt_pts(a.settings.prime_risque_marche, 2, sign=False)}, "
           f"bêta {fmt_num(v.beta, 2)}) ; croissance perpétuelle {fmt_pct(v.g)}.",
           f"Multiples actuels : PER {fmt_x(v.pe)}, P/B {fmt_x(v.pb, 2)}"
           + (f", VE/EBE {fmt_x(v.ev_ebitda)}" if not c.is_bank else "")
           + f" ; médianes des pairs : PER {fmt_x(fd.medians.get('pe'))}, P/B {fmt_x(fd.medians.get('pb'), 2)}.",
           f"Ce que le cours suppose : ROE durable de {fmt_pct(v.implied_roe)} (ROE normatif {fmt_pct(v.roe_norm)})."]
    _bullets(s, Inches(0.6), Inches(1.6) + Inches(0.36 * len(rows)) + Inches(0.2), Inches(5.4), Inches(2.4), hyp,
             size=10.5)
    if v.methods:
        cats = [mm.name for mm in v.methods] + ["Cours actuel"]
        vals = [mm.value for mm in v.methods] + [v.price]
        _bar_chart(s, Inches(6.4), Inches(1.4), Inches(6.4), Inches(5.4), cats, [("FCFA par action", vals)], [S1],
                   fmt='#,##0', title="Valeur par méthode face au cours", horizontal=True, legend=False, gap=50,
                   point_colors={len(cats) - 1: GRAY})

    # 8. Analyse technique ------------------------------------------------------------
    s = d.slide("Analyse technique", f"Tendance {t.trend}" + (f", score {fmt_num(t.score)}/100" if is_num(t.score) else ""),
                notes=tx.get("technique"))
    if t.available:
        img = charts.price_chart(t, months=12, size=(7.6, 4.2), with_rsi=True)
        _image(s, img, Inches(0.5), Inches(1.45), w=Inches(7.6))
        i_, p_, lv = t.indicators, t.performance, t.levels
        rows = [["Indicateur", "Valeur"], ["Moyenne mobile 50 j", fmt_num(i_.get("mm50"))],
                ["Moyenne mobile 200 j", fmt_num(i_.get("mm200"))], ["RSI 14", fmt_num(i_.get("rsi"))],
                ["Volatilité annualisée", fmt_pct(i_.get("volatilite"))], ["Repli max. 12 mois", fmt_pct(i_.get("drawdown_max"))],
                ["Support", fmt_num(lv.get("support_1"))], ["Résistance", fmt_num(lv.get("resistance_1"))],
                ["Perf. relative 6 mois", fmt_pct(p_.get("relative_6m"), sign=True)]]
        _table(s, Inches(8.4), Inches(1.5), Inches(4.4), rows, [2.4, 1.3], font=10.5, row_h=0.31)
        _bullets(s, Inches(8.4), Inches(4.55), Inches(4.4), Inches(2.4), [x for x, _ in t.signals][:5], size=10)
    else:
        _text(s, Inches(0.6), Inches(2), Inches(10), Inches(1), "Historique de cours insuffisant.", size=14, color=MUTED)

    # 9. Macro --------------------------------------------------------------------------
    s = d.slide("Environnement macroéconomique et monétaire", f"{m.country_name} et Union économique et monétaire ouest-africaine",
                notes=tx.get("macro"))
    y0 = m.year or date.today().year
    labels = {"croissance": "Croissance du PIB réel", "inflation": "Inflation moyenne", "dette_publique": "Dette publique / PIB",
              "solde_budgetaire": "Solde budgétaire / PIB", "solde_courant": "Solde courant / PIB"}
    rows = [[f"{m.country_name} (%)", str(y0 - 1), f"{y0}p", f"{y0 + 1}p"]]
    for k, lab in labels.items():
        ser = m.data.get(k, {})
        if ser:
            rows.append([lab] + [fmt_pts(ser.get(yy), 1, sign=False) if is_num(ser.get(yy)) else "n.d."
                                 for yy in (y0 - 1, y0, y0 + 1)])
    _table(s, Inches(0.6), Inches(1.5), Inches(6.2), rows, [2.6, 1, 1, 1], font=11, row_h=0.4)
    b = m.bceao
    _rect(s, Inches(7.2), Inches(1.5), Inches(5.6), Inches(2.3), PAPER, name="Encadré BCEAO")
    _text(s, Inches(7.4), Inches(1.6), Inches(5.3), Inches(0.35), "Politique monétaire BCEAO", size=13, bold=True, color=NAVY)
    _text(s, Inches(7.4), Inches(2.0), Inches(5.3), Inches(1.8),
          [f"Taux directeur : {fmt_pts(b.get('taux_directeur'), 2, sign=False)}",
           f"Guichet de prêt marginal : {fmt_pts(b.get('guichet_pret_marginal'), 2, sign=False)}",
           f"Inflation UEMOA attendue : {fmt_pts(b.get('inflation_uemoa_prevue'), 1, sign=False)} ; "
           f"croissance : {fmt_pts(b.get('croissance_uemoa_prevue'), 1, sign=False)}",
           f"Référence : CPM du {b.get('date_reference')}"], size=11.5)
    _text(s, Inches(7.2), Inches(4.05), Inches(5.6), Inches(0.35), "Sensibilité sectorielle", size=13, bold=True, color=NAVY)
    _text(s, Inches(7.2), Inches(4.45), Inches(5.6), Inches(2.2), m.sector_note, size=11)
    _text(s, Inches(0.6), Inches(4.2), Inches(6.2), Inches(2.5),
          [f"Score macro : {fmt_num(m.score)}/100" if is_num(m.score) else "Score macro : n.d.", f"Source : {m.source}"]
          + m.notes, size=10.5, color=MUTED)

    # 10. Actualités ----------------------------------------------------------------------
    rep = a.interim
    if rep is not None and rep.lines:
        s = d.slide("Exercice en cours : dernière publication",
                    f"{clean_text(rep.title or rep.period)[:110]}"
                    + (f" | publiée le {rep.published:%d/%m/%Y}" if rep.published else ""),
                    notes=" ".join(tx.get("publication", [])))
        rows = [["Indicateur (Mds FCFA)", rep.period or "N", rep.period_prev or "N-1", "Variation"]]
        for ln in rep.lines:
            rows.append([ln.label, fmt_mds(ln.current), fmt_mds(ln.previous),
                         fmt_pct(ln.variation, sign=True) if is_num(ln.variation) else "n.s."])
        _table(s, Inches(0.6), Inches(1.5), Inches(7.0), rows, [3.0, 1.3, 1.3, 1.2], font=10.5,
               row_h=min(0.4, 5.0 / len(rows)))
        main = [ln for ln in rep.lines if ln.key in ("ca", "pnb", "ebe", "rbe", "re", "rao", "rn")][:5]
        if main:
            _bar_chart(s, Inches(7.9), Inches(1.4), Inches(4.9), Inches(3.6), [ln.label for ln in main],
                       [(rep.period_prev or "N-1", [ln.previous / 1e9 if is_num(ln.previous) else None for ln in main]),
                        (rep.period or "N", [ln.current / 1e9 if is_num(ln.current) else None for ln in main])],
                       [GRAY, S1], fmt='#,##0', title="Période en cours face à N-1 (Mds FCFA)", labels_last=False)
        others = [f"{p_.published:%d/%m/%Y} : {clean_text(p_.title)[:95]}" for p_ in a.publications[:4] if p_.published]
        _text(s, Inches(7.9), Inches(5.15), Inches(4.9), Inches(0.3), "Publications récentes", size=12, bold=True,
              color=NAVY)
        _bullets(s, Inches(7.9), Inches(5.5), Inches(4.9), Inches(1.5),
                 others or [f"Source : {clean_text(rep.source or 'document chargé par l analyste')}"], size=9.5)

    s = d.slide("Actualités et informations communiquées", "Veille de presse sur douze mois et communication de la société",
                notes=tx.get("actualites"))
    items = a.news.items[:9]
    if items:
        rows = [["Date", "Source", "Titre", "Ton"]] + [
            [n.published.strftime("%d/%m/%Y") if n.published else "", n.source[:28], n.title[:130],
             "+" if n.sentiment > 0.15 else ("-" if n.sentiment < -0.15 else "=")] for n in items]
        _table(s, Inches(0.6), Inches(1.5), Inches(12.2), rows, [1.1, 1.8, 8.6, 0.5], font=9.5, row_h=0.36,
               align_right_from=3)
        ytxt = Inches(1.6) + Inches(0.36 * len(rows))
    else:
        _text(s, Inches(0.6), Inches(1.5), Inches(12), Inches(0.4), a.news.error or "Aucune actualité recensée.",
              size=12, color=MUTED)
        ytxt = Inches(2.1)
    extra = []
    if a.news.events:
        extra.append("Thèmes relevés : " + ", ".join(f"{k.lower()} ({len(vv)})" for k, vv in a.news.events.items()) + ".")
    if a.news.user_notes:
        extra.append("Informations communiquées par la société : " + a.news.user_notes)
    if tx.get("catalyseurs"):
        extra += ["Catalyseur : " + x for x in tx["catalyseurs"][:3]]
    if extra:
        _bullets(s, Inches(0.6), ytxt, Inches(12.2), Inches(6.9) - ytxt, extra, size=11)

    d.divider(3, "Décision", "Matrice de décision, méthodologie et sources", icon="cible")

    # 11. Décision ---------------------------------------------------------------------------
    s = d.slide("Matrice de décision", f"Signal {dec.signal}, conviction {dec.conviction.lower()}",
                notes=" ".join(tx.get("justification", [])))
    _image(s, charts.decision_matrix_chart(dec, a.settings, size=(4.6, 3.4)), Inches(0.5), Inches(1.4), h=Inches(4.1))
    keys = list(config.PILLAR_LABELS)
    _bar_chart(s, Inches(6.0), Inches(1.35), Inches(6.8), Inches(2.9), [config.PILLAR_LABELS[k] for k in keys],
               [("Score", [dec.pillars.get(k) for k in keys])], [S1], fmt='0', title="Notation par pilier (/100)",
               horizontal=True, legend=False, gap=50)
    just = list(tx.get("justification", []))[:-1] + dec.warnings
    _bullets(s, Inches(6.0), Inches(4.35), Inches(6.8), Inches(2.6), just[:5], size=10.5)
    _text(s, Inches(0.6), Inches(5.65), Inches(5.2), Inches(1.3), config.SIGNAL_DEFINITIONS.get(dec.signal, ""),
          size=10.5, italic=True, color=MUTED)

    # 12. Méthodologie --------------------------------------------------------------------------
    s = d.slide("Méthodologie, sources et avertissement")
    _text(s, Inches(0.6), Inches(1.4), Inches(7.2), Inches(0.35), "Méthode", size=13, bold=True, color=NAVY)
    _text(s, Inches(0.6), Inches(1.8), Inches(7.2), Inches(3.2), METHODO, size=10.5)
    src = dict(a.market.sources)
    if c.level in ("complet", "partiel"):
        src["états financiers"] = (f"Classeur, feuille {c.sheet} ({c.unit_label})" if c.sheet
                                   else f"{c.data_source} ({c.unit_label})")
    else:
        src["chiffres clés"] = c.data_source
    src["macroéconomie"] = m.source
    _text(s, Inches(8.2), Inches(1.4), Inches(4.6), Inches(0.35), "Sources", size=13, bold=True, color=NAVY)
    _bullets(s, Inches(8.2), Inches(1.8), Inches(4.6), Inches(3.4), [f"{k.capitalize()} : {vv}" for k, vv in src.items()],
             size=9.5, space_after=2)
    notes_dq = c.notes[:3]
    if notes_dq:
        _text(s, Inches(0.6), Inches(5.0), Inches(12.2), Inches(0.9), ["Contrôle des données : " + x for x in notes_dq],
              size=9.5, color=MUTED)
    _text(s, Inches(0.6), Inches(6.05), Inches(12.2), Inches(0.9), config.DISCLAIMER, size=9.5, italic=True, color=MUTED)

    d.close()

    buf = io.BytesIO()
    d.prs.core_properties.title = clean_text(f"Fiche de valeur {c.name} ({c.ticker})")
    d.prs.core_properties.author = clean_text(a.settings.maison)
    d.prs.save(buf)
    return buf.getvalue()


def _price_native(slide, t, x, y, w, h, months=18):
    import pandas as pd
    df = t.frame
    df = df[df.index >= df.index[-1] - pd.DateOffset(months=months)]
    if len(df) > 260:
        df = df.iloc[:: max(1, len(df) // 260)]
    cats = [ix.date() for ix in df.index]
    series = [("Cours de clôture", list(df["close"])), ("Moyenne mobile 50 j", list(df["mm50"])),
              ("Moyenne mobile 200 j", list(df["mm200"]))]
    ch = _line_chart(slide, x, y, w, h, cats, series, [S1, S2, S3], fmt='#,##0',
                     title=f"Cours sur {months} mois (FCFA)", widths=[2.0, 1.5, 1.5], date_axis=True)
    # Échelle des cours resserrée autour des données (pas de départ à zéro pour un cours)
    import math
    vals = [v for col in ("close", "mm50", "mm200") for v in df[col].tolist() if is_num(v)]
    if vals:
        lo, hi = min(vals), max(vals)
        step = 10 ** max(0, int(math.log10(max(hi - lo, 1))))
        ch.value_axis.minimum_scale = max(0, math.floor(lo * 0.97 / step) * step)
        ch.value_axis.maximum_scale = math.ceil(hi * 1.02 / step) * step
    # Axe des dates : une graduation tous les trois mois
    from lxml import etree
    from pptx.oxml.ns import qn
    plot_area = ch._chartSpace.chart.plotArea
    date_ax = plot_area.find(qn("c:dateAx"))
    if date_ax is not None:
        for tag in ("c:baseTimeUnit", "c:majorUnit", "c:majorTimeUnit"):
            for el in date_ax.findall(qn(tag)):
                date_ax.remove(el)
        ext = date_ax.find(qn("c:extLst"))
        for tag, val in (("c:baseTimeUnit", "days"), ("c:majorUnit", "3"), ("c:majorTimeUnit", "months")):
            el = etree.SubElement(date_ax, qn(tag))
            el.set("val", val)
            if ext is not None:
                ext.addprevious(el)
    return ch
