"""Tests hors ligne : lecture du classeur, analyses, rapports et analyseurs de flux.

Lancement : python -m pytest -q
"""
from __future__ import annotations

import io
import json
import re
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from fiche_valeur import config
from fiche_valeur.config import Settings
from fiche_valeur.market import load_market_base, parse_capitalisations, parse_sika_payload
from fiche_valeur.news import parse_rss
from fiche_valeur.parser import parse_workbook
from fiche_valeur.pipeline import run_analysis
from fiche_valeur.utils import clean_text, fmt_num, fmt_pct, norm_label_core

ROOT = Path(__file__).resolve().parents[1]
EF = ROOT / "sample_data" / "demo_etats_financiers.xlsx"
SMF = ROOT / "sample_data" / "demo_base_marche.xlsx"
SIGNALS = set(config.SIGNAL_COLORS)


@pytest.fixture(scope="module")
def wb():
    return parse_workbook(EF)


@pytest.fixture(scope="module")
def mb():
    return load_market_base(SMF)


@pytest.fixture(scope="module")
def offline():
    return Settings(en_ligne=False)


@pytest.fixture(scope="module")
def analyses(wb, mb, offline):
    return {t: run_analysis(wb, t, offline, market_base=mb) for t in wb.tickers}


# --------------------------------------------------------------------------- formats

def test_percent_attached_and_no_long_dash():
    assert fmt_pct(0.125) == "12,5%"
    assert " %" not in fmt_pct(-0.034)
    assert clean_text("Marge — en hausse – 10 %") == "Marge - en hausse - 10%"
    assert fmt_num(1234567.8, 0).replace(" ", " ").replace("\xa0", " ") == "1 234 568"


def test_label_normalisation():
    assert norm_label_core("XB - Chiffre d'affaires (A+B)") == "CHIFFRE D AFFAIRES"
    assert norm_label_core("Résultat net") == "RESULTAT NET"


# --------------------------------------------------------------------------- lecture

def test_demo_workbook(wb):
    assert set(wb.tickers) == {"DBQC", "DBQS", "DAGC", "DDIC", "DTLS", "DENC"}
    kinds = {t: c.kind for t, c in wb.companies.items()}
    assert kinds["DBQC"] == "banque" and kinds["DAGC"] == "syscohada" and kinds["DTLS"] == "ifrs"
    for c in wb.companies.values():
        assert c.years[-1] == 2025
        assert c.get(c.net_income_field, 2025) is not None
        assert c.get(c.equity_field, 2025) is not None


def test_unit_harmonisation(wb):
    c = wb.companies["DDIC"]
    assert any("2021" in n and "unité" in n for n in c.notes)
    ca = [c.get("ca", y) for y in c.years]
    # Après harmonisation, aucun exercice ne s'écarte d'un facteur 10 de la médiane
    med = sorted(ca)[len(ca) // 2]
    assert all(0.1 < v / med < 10 for v in ca)


def test_formula_evaluation(tmp_path):
    """Classeur saisi avec formules non recalculées (pas de valeurs en cache)."""
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "TEST"
    ws.append(["Test Industries SA - états financiers (en millions FCFA)"])
    ws.append(["Poste", 2023, 2024, 2025])
    ws.append(["COMPTE DE RÉSULTAT"])
    ws.append(["Chiffre d'affaires", 50000, 55000, 60500])
    ws.append(["Achats consommés", -30000, -32000, -35000])
    ws.append(["Valeur ajoutée", "=B4+B5", "=C4+C5", "=D4+D5"])
    ws.append(["Charges de personnel", -6000, -6500, -7000])
    ws.append(["Excédent brut d'exploitation", "=B6+B7", "=C6+C7", "=D6+D7"])
    ws.append(["Dotations aux amortissements", -3000, -3200, -3400])
    ws.append(["Résultat d'exploitation", "=SUM(B8:B9)", "=SUM(C8:C9)", "=SUM(D8:D9)"])
    ws.append(["Résultat net", 7000, 8100, 9400])
    ws.append(["BILAN"])
    ws.append(["Total actif", 90000, 95000, 101000])
    ws.append(["Capitaux propres", 40000, 44000, 48000])
    ws.append(["Dettes financières", 15000, 14000, 13000])
    path = tmp_path / "formules.xlsx"
    book.save(path)
    c = parse_workbook(path).companies["TEST"]
    assert c.kind == "syscohada"
    assert c.get("va", 2025) == pytest.approx(25500e6)
    assert c.get("ebe", 2025) == pytest.approx(18500e6)
    assert c.get("re", 2024) == pytest.approx(13300e6)


def test_market_base(mb):
    assert "DBQC" in mb.tickers
    assert mb.closes is not None and len(mb.closes) > 200
    assert mb.shares_total.get("DBQC")


# --------------------------------------------------------------------------- analyses

def test_signals_offline(analyses):
    for t, a in analyses.items():
        d = a.decision
        assert d.signal in SIGNALS, t
        assert 0 <= d.composite <= 100
        assert d.trace and "signal de base" in d.trace[0]
        assert a.valuation.fair_value and a.valuation.fair_value > 0
        assert a.technical.available
        for txt in a.texts.values():
            flat = " ".join(txt) if isinstance(txt, list) else str(txt)
            assert "—" not in flat, t
            assert not re.search(r"\d[ \xa0\u202f]%", flat), (t, flat)


def test_signal_diversity(analyses):
    signals = {a.decision.signal for a in analyses.values()}
    assert "ACHAT" in signals and "VENDRE" in signals


def test_major_event_forces_wait(wb, mb, offline):
    a = run_analysis(wb, "DAGC", offline, market_base=mb, major_event=True)
    assert a.decision.signal in {"ATTENDRE", "VENDRE"}


def test_manual_price_override(wb, mb, offline):
    base = run_analysis(wb, "DAGC", offline, market_base=mb)
    cheap = run_analysis(wb, "DAGC", offline, market_base=mb, overrides={"price": base.valuation.price * 0.6})
    assert cheap.valuation.upside > base.valuation.upside


# --------------------------------------------------------------------------- rapports

def test_reports(analyses):
    from pptx import Presentation

    from fiche_valeur.export import build_excel, build_screening_excel, screening_frame
    from fiche_valeur.report_pdf import build_pdf
    from fiche_valeur.report_pptx import build_pptx

    a = analyses["DBQC"]
    pdf = build_pdf(a)
    assert pdf[:4] == b"%PDF" and len(pdf) > 50_000
    assert build_pdf(a, compact=True)[:4] == b"%PDF"
    prs = Presentation(io.BytesIO(build_pptx(analyses["DAGC"])))
    assert len(prs.slides) >= 10
    xl = pd.read_excel(io.BytesIO(build_excel(a)), sheet_name=None)
    assert {"Synthèse", "SIG", "Ratios", "Valorisation"} <= set(xl)
    scr = screening_frame(list(analyses.values()))
    assert len(scr) == len(analyses)
    assert build_screening_excel(scr)[:2] == b"PK"


# --------------------------------------------------------------------------- flux en ligne (hors ligne)

CAPI_HTML = """
<table><thead><tr><th>Symbole</th><th>Nom</th><th>Nombre de titres</th><th>Cours du jour</th>
<th>Capitalisation flottante</th><th>Capitalisation globale</th><th>Capitalisation globale (%)</th></tr></thead>
<tbody>
<tr><td>SNTS</td><td>SONATEL SN</td><td>100 000 000</td><td>25 500</td><td>675 750 000 000</td>
<td>2 550 000 000 000</td><td>20,1</td></tr>
<tr><td>SGBC</td><td>SOCIETE GENERALE COTE D'IVOIRE</td><td>31 111 110</td><td>28 000</td>
<td>145 000 000 000</td><td>871 111 080 000</td><td>6,9</td></tr>
</tbody></table>"""


def test_parse_capitalisations():
    df = parse_capitalisations(CAPI_HTML)
    assert list(df.index) == ["SNTS", "SGBC"]
    assert df.loc["SNTS", "nb_titres"] == 100_000_000
    assert df.loc["SGBC", "cours"] == 28_000
    assert df.loc["SNTS", "cap_globale"] == 2.55e12


def test_parse_sika_payload():
    payload = json.loads("""{"lst": [
        {"Date": "02/01/2025", "Open": 100, "High": 104, "Low": 99, "Close": 103, "Volume": 1200},
        {"Date": "03/01/2025", "Open": 103, "High": 106, "Low": 101, "Close": 105, "Volume": 800},
        {"Date": "06/01/2025", "Open": 105, "High": 105, "Low": 98, "Close": 99, "Volume": 1500}]}""")
    df = parse_sika_payload(payload)
    assert len(df) == 3 and df["close"].iloc[-1] == 99
    assert df.index.is_monotonic_increasing
    assert parse_sika_payload({"lst": []}) is None


RSS = """<?xml version="1.0" encoding="UTF-8"?><rss version="2.0"><channel>
<item><title>Sonatel : le chiffre d'affaires progresse de 8% - Sika Finance</title><link>https://ex.org/1</link>
<pubDate>Mon, 05 Oct 2026 08:00:00 GMT</pubDate><source url="https://ex.org">Sika Finance</source></item>
<item><title>Sonatel propose un dividende record aux actionnaires - Agence Ecofin</title><link>https://ex.org/2</link>
<pubDate>Tue, 15 Sep 2026 08:00:00 GMT</pubDate><source url="https://ex.org">Agence Ecofin</source></item>
<item><title>Le marché régional des télécoms en Afrique de l'Ouest - Jeune Afrique</title><link>https://ex.org/3</link>
<pubDate>Tue, 15 Sep 2026 08:00:00 GMT</pubDate><source url="https://ex.org">Jeune Afrique</source></item>
</channel></rss>"""


def test_parse_rss():
    news = parse_rss(RSS.encode(), "SNTS", "Sonatel")
    titles = [n.title for n in news.items]
    assert len(titles) == 2 and all("Sonatel" in t for t in titles)
    assert not titles[0].endswith("Sika Finance")
    assert "Dividende / détachement" in news.events
    assert news.sentiment > 0 and news.score > 50
    assert parse_rss(b"pas du xml", "SNTS", "Sonatel").error


# --------------------------------------------------------------------------- titres à données partielles

SIKA_SOCIETE_HTML = """<html><body><div>Les chiffres sont en millions de FCFA</div>
<table><tr><td></td><td>2021</td><td>2022</td><td>2023</td><td>2024</td><td>2025</td></tr>
<tr><td>Chiffre d'affaires</td><td>80 000</td><td>88 000</td><td>95 000</td><td>104 000</td><td>112 000</td></tr>
<tr><td>Croissance CA</td><td></td><td>10,00 %</td><td>7,95 %</td><td>9,47 %</td><td>7,69 %</td></tr>
<tr><td>Résultat net</td><td>18 000</td><td>20 000</td><td>22 000</td><td>24 000</td><td>27 000</td></tr>
<tr><td>BNPA</td><td>720,00</td><td>800,00</td><td>880,00</td><td>960,00</td><td>1 080,00</td></tr>
<tr><td>PER</td><td>10,5</td><td>9,8</td><td>9,1</td><td>8,9</td><td>8,6</td></tr>
<tr><td>Dividende</td><td>400,00</td><td>450,00</td><td>500,00</td><td>550,00</td><td>600,00</td></tr></table>
<p>Nombre de titres : 25 000 000</p><p>Flottant : 22,5 %</p>
<p>Principaux actionnaires : GROUPE X 60 % ; PUBLIC (BRVM) 22,5 %</p></body></html>"""


def test_parse_sika_societe():
    from fiche_valeur.partial import parse_sika_societe
    kf = parse_sika_societe(SIKA_SOCIETE_HTML)
    assert kf.revenue[2025] == 112_000e6 and kf.net_income[2021] == 18_000e6
    assert kf.dps[2023] == 500 and kf.eps[2025] == 1080 and kf.per[2021] == 10.5
    assert kf.shares == 25_000_000 and kf.float_pct == pytest.approx(0.225)
    assert parse_sika_societe("<html><body>rien</body></html>") is None


def test_manual_figures_table():
    from fiche_valeur.partial import empty_manual_table, figures_from_table
    assert figures_from_table(empty_manual_table()) is None
    df = empty_manual_table(2025, 3)
    df.iloc[:, 1] = [100.0, 110.0, 120.0]       # chiffre d'affaires, millions
    df.iloc[:, 2] = [10.0, 12.0, 15.0]          # résultat net
    kf = figures_from_table(df)
    assert kf.revenue[2025] == 120e6 and kf.net_income[2023] == 10e6


@pytest.fixture(scope="module")
def partial_env(mb):
    config.register_ticker("DAGC", "Agro démo", "CI", "Consommation de base")
    config.register_ticker("DBQC", "Banque démo", "CI", "Services financiers")
    return mb


def test_level2_key_figures(partial_env, offline):
    from fiche_valeur.partial import parse_sika_societe
    kf = parse_sika_societe(SIKA_SOCIETE_HTML)
    a = run_analysis(None, "DAGC", offline, market_base=partial_env, figures=kf)
    assert a.company.level == "chiffres_cles"
    assert a.decision.signal in SIGNALS and a.decision.conviction in {"Moyenne", "Faible"}
    assert a.valuation.pe and a.valuation.fair_value
    assert a.coverage.level == "chiffres_cles" and 0 < a.coverage.score < 100
    assert any(i.name.startswith("Bilan") and i.status == "absent" for i in a.coverage.items)


def test_level3_market_only(partial_env, offline):
    a = run_analysis(None, "DBQC", offline, market_base=partial_env)
    assert a.company.level == "marche"
    d = a.decision
    assert d.signal in {"ATTENDRE", "CONSERVER", "ALLÉGER", "ACHAT"} and d.signal != "VENDRE"
    assert d.conviction == "Faible" and d.fundamental is None
    assert a.technical.available and a.valuation.fair_value


def test_level3_without_any_data(offline):
    config.register_ticker("ZZZZ", "Titre sans données", "CI", "Non classé")
    a = run_analysis(None, "ZZZZ", offline)
    assert a.decision.signal == "ATTENDRE" and a.decision.conviction == "Faible"
    assert a.coverage.score < 40


def test_partial_reports(partial_env, offline):
    from pptx import Presentation

    from fiche_valeur.export import build_excel
    from fiche_valeur.partial import parse_sika_societe
    from fiche_valeur.report_pdf import build_pdf
    from fiche_valeur.report_pptx import build_pptx
    cases = [run_analysis(None, "DAGC", offline, market_base=partial_env, figures=parse_sika_societe(SIKA_SOCIETE_HTML)),
             run_analysis(None, "DBQC", offline, market_base=partial_env),
             run_analysis(None, "ZZZZ", offline)]
    for a in cases:
        assert build_pdf(a)[:4] == b"%PDF" and build_pdf(a, compact=True)[:4] == b"%PDF"
        assert len(Presentation(io.BytesIO(build_pptx(a))).slides) >= 9
        assert "Couverture" in pd.read_excel(io.BytesIO(build_excel(a)), sheet_name=None)


def test_guess_country():
    assert config.guess_country("BOABF") == "BF" and config.guess_country("BOAS") == "SN"
    assert config.guess_country("SGBC") == "CI" and config.guess_country("ETIT") == "TG"


def test_full_referential():
    assert len(config.TICKERS) >= 47
    assert all(v["pays"] in config.COUNTRIES and v["secteur"] in config.SECTOR_MULTIPLES
               for v in config.TICKERS.values())
